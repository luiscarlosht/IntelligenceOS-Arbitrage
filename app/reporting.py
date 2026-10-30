"""Formatted, formula-safe Excel reports written atomically."""
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ['source', 'product', 'upc', 'amazon_asin', 'destination', 'buy_cost',
           'sale_price', 'profit', 'roi', 'status', 'reason', 'source_url']

DISCOVERY_COLUMNS = COLUMNS + [
    'manufacturer', 'model', 'package_quantity', 'variation', 'availability',
    'referral_fee', 'fulfillment_fee', 'inbound_shipping', 'preparation', 'packaging',
    'storage', 'sales_tax', 'other_cost', 'transaction_cost', 'acquisition_cost',
    'break_even_price', 'maximum_purchase_price', 'monthly_sales', 'offer_count',
    'matching_confidence', 'confidence', 'ranking_score', 'budget_units', 'cost_basis',
    'unverified_costs', 'source_timestamp', 'marketplace_timestamp', 'source_provenance',
    'marketplace_provenance', 'data_mode', 'purchase_ready',
]


def write_report(rows, destination, *, discovery=False, top_count=20, job=None, demo=False,
                 enrichment_rows=None, candidate_review=None):
    if candidate_review is not None and (not discovery or demo):
        raise ValueError('Candidate review requires discovery and cannot relabel DEMO evidence')
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = 'Executive Summary'
    for row in [
        ['Proof of concept — candidate review' if candidate_review else ('DEMO — FICTIONAL FIXTURES ONLY' if demo else 'INTELLIGENCEOS ARBITRAGE'),
         'Saved observations — identity and profitability unverified; not a purchase recommendation' if candidate_review
         else ('MOCK / supplied analysis — never purchase approval' if discovery else 'Illustrative sourcing analysis')],
        ['Run (UTC)', datetime.now(timezone.utc).isoformat()],
        ['Candidates', len(rows)],
        ['Review candidates', sum(r['status'] == 'REVIEW' for r in rows)],
        ['Rejected', sum(r['status'] == 'REJECT' for r in rows)],
        ['Errors', sum(r['status'] == 'ERROR' for r in rows)],
        ['Notice', 'Verify live prices, eligibility, fees and exact matches before purchase.'],
    ]:
        summary.append(row)
    if candidate_review:
        for key in ('Source evidence', 'Amazon evidence', 'Economics', 'Confirmed exact matches',
                    'Candidate-only rows', 'Unmatched rows'):
            summary.append([key, candidate_review[key]])
    columns = DISCOVERY_COLUMNS if discovery else COLUMNS
    selected_top = sorted((r for r in rows if r['status'] == 'REVIEW'),
                          key=lambda r: r['ranking_score'] if discovery else r['roi'], reverse=True)
    if discovery: selected_top = selected_top[:top_count]
    for name, selected in [
        ('Top 20 Opportunities' if discovery and top_count == 20 else ('Top Opportunities' if not discovery else 'Top Ranked Opportunities'), selected_top),
        ('All Candidates', rows),
        ('Rejected Products', [r for r in rows if r['status'] != 'REVIEW']),
    ]:
        sheet = wb.create_sheet(name)
        sheet.append([c.replace('_', ' ').title() for c in columns])
        for result in selected:
            values = [result.get(c, '') for c in columns]
            for i in (5, 6):
                try:
                    values[i] = float(values[i])
                    import math
                    if not math.isfinite(values[i]): values[i] = str(result.get(COLUMNS[i], ''))
                except (ValueError, TypeError):
                    pass
            sheet.append(values)
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for i in (5, 6, 7): row[i].number_format = '$#,##0.00'
            row[8].number_format = '0.0"%"'
            if discovery:
                for index, key in enumerate(columns):
                    if key in ('referral_fee', 'fulfillment_fee', 'inbound_shipping', 'preparation', 'packaging',
                               'storage', 'sales_tax', 'other_cost', 'transaction_cost', 'acquisition_cost',
                               'break_even_price', 'maximum_purchase_price'):
                        row[index].number_format = '$#,##0.00'
                url = row[11].value
                from app.normalization import http_url
                if isinstance(url, str) and http_url(url): row[11].hyperlink = url
            row[9].font = Font(color={'REVIEW': '137333', 'REJECT': 'B06000', 'ERROR': 'B00020'}.get(row[9].value, '000000'), bold=True)
    if discovery:
        quality = wb.create_sheet('Data Quality and Missing Info')
        quality.append(['Source', 'Product', 'ASIN', 'Issues', 'Unverified Costs', 'Data Mode', 'Purchase Ready'])
        for result in rows:
            quality.append([result.get('source'), result.get('product'), result.get('amazon_asin'),
                            result.get('quality_issues') or 'No missing critical screening data',
                            result.get('unverified_costs'), result.get('data_mode'), False])
        quality.auto_filter.ref = quality.dimensions
        execution = wb.create_sheet('Job Execution Summary')
        execution.append(['Field', 'Value'])
        for key, value in (job or {'phase': 'Standalone report; final status unavailable'}).items():
            execution.append([key, value])
        execution.append(['Report generated UTC', datetime.now(timezone.utc).isoformat()])
        execution.append(['Warning', 'Saved evidence only; no new API calls. API success does not establish identity, economics or email delivery.'
                          if candidate_review else 'Mock data is fictional. SQLite is authoritative for final job and delivery outcome.'])
    if demo:
        wb.properties.title = 'DEMO — fictional fixtures; never live deals'
        for sheet, title in zip(wb, ('DEMO Summary', 'DEMO Review Candidates', 'DEMO All Candidates',
                                     'DEMO Rejections', 'DEMO Data Quality', 'DEMO Execution')):
            sheet.title = title
    if enrichment_rows is not None:
        evidence = wb.create_sheet('DEMO Enrichment' if demo else 'Enrichment Evidence')
        fields = ('source_observation_id', 'asin', 'kind', 'value', 'observed_at',
                  'fetched_at', 'provenance', 'is_mock', 'fresh', 'notice')
        evidence.append([field.replace('_', ' ').title() for field in fields])
        for result in enrichment_rows:
            evidence.append([result.get(field) for field in fields])
        evidence.auto_filter.ref = evidence.dimensions
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='12304A')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet:
            for cell in row:
                # Force all user strings to literal cells, including spreadsheet formulas.
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 55 if i in (2, 11, 12) else 21
    if candidate_review:
        sheet = wb.create_sheet('Candidate Comparison', 1)
        headers = ('Source Title', 'Seller', 'Observed Item Price (USD)', 'Reported Shipping (USD)',
                   'Source URL', 'Candidate ASIN', 'Amazon Title', 'Match Status', 'Evidence Gaps')
        from app.candidate_enrichment_report import COLUMNS as AMAZON_COLUMNS
        sheet.append(headers + tuple(c[1] for c in AMAZON_COLUMNS))
        for index, result in enumerate(candidate_review['comparisons'], 2):
            sheet.append([result.get(key) for key in ('ebay_title', 'seller', 'observed_item_price',
                          'reported_shipping', 'source_url', 'candidate_asin', 'amazon_title', 'match_status', 'evidence_gaps')]
                         + [result.get(c[0]) for c in AMAZON_COLUMNS])
            sheet.row_dimensions[index].height = 120
            for col in (3, 4):
                sheet.cell(index, col).number_format = '$#,##0.00'
            from app.normalization import http_url
            url = result.get('source_url')
            if isinstance(url, str) and http_url(url):
                # Display a readable listing URL; the hyperlink retains the full supplied URL.
                from urllib.parse import urlsplit, urlunsplit
                parsed = urlsplit(url)
                sheet.cell(index, 5, urlunsplit((parsed.scheme, parsed.netloc, parsed.path, '', '')))
                sheet.cell(index, 5).hyperlink = url
                sheet.cell(index, 5).font = Font(color='0563C1', underline='single')
        from app.amazon_approval_report import write_approval_links
        write_approval_links(wb, candidate_review['comparisons'])
        for offset, column in enumerate(AMAZON_COLUMNS, 10):
            sheet.column_dimensions[get_column_letter(offset)].width = column[2]
            if column[3]:
                for row in range(2, sheet.max_row+1):
                    sheet.cell(row, offset).number_format = '$#,##0.00'
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 80
        for col, width in enumerate((48, 23, 19, 19, 32, 17, 60, 23, 85), 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='12304A')
            cell.font = Font(color='FFFFFF', bold=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.print_title_rows = '1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        wb.properties.title = 'Proof of concept — candidate review'
    temp = None
    try:
        with NamedTemporaryFile(dir=destination.parent, suffix='.xlsx', delete=False) as f:
            temp = Path(f.name)
        wb.save(temp)
        temp.replace(destination)
    finally:
        wb.close()
        if temp: temp.unlink(missing_ok=True)
