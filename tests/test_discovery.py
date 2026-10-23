"""V0.2 contracts and end-to-end fixtures; no live API or SMTP requests."""
import csv
import json
import sqlite3
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook
from app.config import Config, ROOT
from app.connectors import CsvConnector
from app.discovery import discover, screen
from app.marketplace import MockAmazonProvider, match_listings
from app.models import RetailProduct, MarketplaceListing
from app.normalization import gtin, money, timestamp, integer, http_url
from app.orchestrator import run_job
from app.profitability import analyze, COST_FIELDS
from app.reporting import write_report

NOW = datetime(2026, 10, 23, 12, tzinfo=timezone.utc)

@pytest.fixture
def config(tmp_path):
    return Config(ROOT/'sample_data/discovery.csv', tmp_path/'report.xlsx', tmp_path/'history.db', mode='discovery')

@pytest.fixture
def product():
    return RetailProduct('Walmart (MOCK)', 'Widget', 'https://example.org/item', gtin('036000291452'),
                         'Acme', 'A-1', 1, D('8'), 'in_stock', NOW, 'fictional supplied CSV', variation='blue')

@pytest.fixture
def listing(product):
    costs = dict.fromkeys(COST_FIELDS, D('0'))
    costs.update(referral_fee=D('3.75'), fulfillment_fee=D('4.5'), inbound_shipping=D('1.25'), other_cost=D('.5'))
    return MarketplaceListing('B000000001', product.gtin, product.manufacturer, product.model,
        1, 'blue', True, D('25'), costs, monthly_sales=60, offer_count=4,
        retrieved_at=NOW, provenance='mock quote', verified_costs=())


def screened(product, listing, config):
    return screen(product, match_listings(product, (listing,)), config, NOW)

@pytest.mark.parametrize('raw', ['036000291452', '0036000291452', '00036000291452'])
def test_gtin_normalized(raw):
    assert gtin(raw) == '00036000291452'

@pytest.mark.parametrize('raw', ['036000291453', '123', 'abcdefghijkl', '０３６０００２９１４５２'])
def test_invalid_gtin(raw):
    with pytest.raises(ValueError): gtin(raw)

@pytest.mark.parametrize('raw', ['NaN', 'Infinity', '-1', 'one'])
def test_invalid_money(raw):
    with pytest.raises(ValueError): money(raw)


def test_timestamps_and_numbers():
    assert timestamp('2026-10-23T12:00:00Z') == NOW
    assert money('') is None
    assert integer('') is None
    with pytest.raises(ValueError): timestamp('2026-10-23T12:00:00')
    with pytest.raises(ValueError): integer('1.2')
    assert not http_url('https://user:password@example.org')


def test_csv_normalizes_and_preserves_bad_rows(tmp_path):
    path = tmp_path/'products.csv'
    path.write_text('retailer,name,purchase_price,gtin,package_quantity,retrieved_at\nTest,Good,8,036000291452,1,2026-10-23T12:00:00Z\nTest,Bad,nan,123,0,wrong\n')
    good, bad = list(CsvConnector(path).products())
    assert good.gtin == '00036000291452'
    assert good.purchase_price == D('8')
    assert bad.purchase_price is None
    assert len(bad.issues) == 4

@pytest.mark.parametrize('content', ['product\nwrong\n', 'retailer,name,purchase_price,name\na,b,1,c\n'])
def test_csv_bad_headers(tmp_path, content):
    path = tmp_path/'bad.csv'; path.write_text(content)
    with pytest.raises(ValueError): list(CsvConnector(path).products())


def test_exact_gtin_and_model_matching(product, listing):
    assert match_listings(product, (listing,)).exact
    assert match_listings(replace(product, gtin=''), (listing,)).exact
    # GTIN alone never establishes seller eligibility.
    result = match_listings(product, (replace(listing, seller_eligible=False),))
    assert result.exact and result.listing.seller_eligible is False

@pytest.mark.parametrize('change', [dict(package_quantity=2), dict(package_quantity=None),
    dict(variation='red'), dict(variation=''), dict(manufacturer='Other'), dict(model='A1')])
def test_matching_mismatches(product, listing, change):
    assert not match_listings(product, (replace(listing, **change),)).exact


def test_asin_hint_needs_identity(product, listing):
    product = replace(product, gtin='', manufacturer='', model='', asin_hint=listing.asin)
    assert not match_listings(product, (listing,)).exact


def test_asin_hint_conflicting_gtin(product, listing):
    product = replace(product, asin_hint=listing.asin)
    assert not match_listings(product, (replace(listing, gtin=gtin('012345678905')),)).exact


def test_ambiguous_and_absent_matches(product, listing):
    assert not match_listings(product, (listing, replace(listing, asin='B000000002'))).exact
    assert match_listings(product, ()).listing is None


def test_costs_preserve_v01(listing):
    result = analyze(D('8'), D('25'), listing.costs)
    assert result.net_profit == D('7')
    assert round(result.roi, 1) == D('71.8')
    assert result.break_even_price == D('18')
    assert round(result.maximum_purchase_price, 2) == D('9.42')


def test_extended_costs_and_rate(listing):
    costs = dict(listing.costs, preparation=D('.5'), packaging=D('.25'), storage=D('.1'), sales_tax=D('.6'), transaction_cost=D('.2'))
    result = analyze(D('8'), D('25'), costs, referral_rate=D('.15'))
    assert result.acquisition_cost == D('11.1')
    assert result.net_profit == D('5.35')
    assert result.break_even_price == D('15.9') / D('.85')
    # Maximum purchase price satisfies both constraints, including tax and prep.
    at_limit = analyze(result.maximum_purchase_price, D('25'), costs, referral_rate=D('.15'))
    assert abs(at_limit.roi - 50) < D('.0000001')
    assert at_limit.net_profit >= 5
    above_limit = analyze(result.maximum_purchase_price + D('.01'), D('25'), costs, referral_rate=D('.15'))
    assert above_limit.roi < 50


def test_exact_profit_roi_boundaries(listing):
    costs = dict.fromkeys(COST_FIELDS, D('0'))
    e = analyze(D('10'), D('15'), costs)
    assert e.net_profit == 5 and e.roi == 50 and e.maximum_purchase_price == 10

@pytest.mark.parametrize('purchase,sale,rate', [('0','25',None), ('8','nan',None), ('8','25','1'), ('8','25','-0.1')])
def test_invalid_economics(listing, purchase, sale, rate):
    with pytest.raises(ValueError): analyze(D(purchase), D(sale), listing.costs, referral_rate=D(rate) if rate else None)


def test_mock_never_purchase_ready(product, listing, config):
    row = screened(product, listing, config)
    assert row['status'] == 'REVIEW'
    assert row['purchase_ready'] is False
    assert row['confidence'] <= 50
    assert row['cost_basis'] == 'estimated'
    assert row['data_mode'].startswith('MOCK')
    assert row['budget_units'] == 51

@pytest.mark.parametrize('change,reason', [
    (dict(seller_eligible=None), 'Seller eligibility'),
    (dict(monthly_sales=None), 'Missing demand'),
    (dict(monthly_sales=29), 'Demand below'),
    (dict(offer_count=16), 'Competition above'),
    (dict(offer_count=None), 'Missing competition'),
    (dict(retrieved_at=NOW-timedelta(hours=25)), 'Stale marketplace'),
    (dict(retrieved_at=NOW+timedelta(seconds=1)), 'Future marketplace'),
    (dict(sale_price=None), 'Missing sale price'),
    (dict(provenance=''), 'Missing marketplace provenance'),
])
def test_marketplace_gates(product, listing, config, change, reason):
    row = screened(product, replace(listing, **change), config)
    assert row['status'] == 'REJECT'
    assert reason in row['quality_issues']

@pytest.mark.parametrize('change,reason', [
    (dict(availability='unknown'), 'Inventory not confirmed'),
    (dict(retrieved_at=None), 'Missing source'),
    (dict(provenance=''), 'Missing source provenance'),
    (dict(url=''), 'Missing source url'),
    (dict(purchase_price=None), 'Missing source purchase_price'),
])
def test_source_gates(product, listing, config, change, reason):
    row = screened(replace(product, **change), listing, config)
    assert row['status'] == 'REJECT'
    assert reason in row['quality_issues']


def test_missing_cost_is_not_zero(product, listing, config):
    costs = dict(listing.costs); del costs['storage']
    row = screened(product, replace(listing, costs=costs), config)
    assert row['status'] == 'REJECT'
    assert row['profit'] is None
    assert 'Missing cost storage' in row['quality_issues']


def test_rate_without_fixed_referral_fee(product, listing, config):
    costs = dict(listing.costs); del costs['referral_fee']
    row = screened(product, replace(listing, costs=costs, referral_rate=D('.15')), config)
    assert row['profit'] == 7
    assert 'referral_rate' in row['unverified_costs']


def test_budget_gate(product, listing, config):
    row = screened(product, listing, replace(config, inventory_budget=5))
    assert row['status'] == 'REJECT'
    assert 'budget' in row['reason']


def test_fixture_forces_mock_even_if_claimed_live(tmp_path):
    path = tmp_path/'quotes.json'
    path.write_text(json.dumps([dict(asin='A', is_mock=False)]))
    assert MockAmazonProvider(path).listings[0].is_mock is True
    path.write_text(json.dumps([dict(asin='A', seller_eligible='yes')]))
    with pytest.raises(ValueError): MockAmazonProvider(path)


def test_discovery_ranking_and_limits(product, listing, config):
    class Connector:
        def products(self): return [product, replace(product, name='Lower ROI', purchase_price=D('9'))]
    class Provider:
        def match(self, p): return match_listings(p, (listing,))
    rows = discover(Connector(), Provider(), config, NOW)
    assert rows[0]['product'] == 'Widget'
    assert rows[0]['ranking_score'] > rows[1]['ranking_score']
    with pytest.raises(ValueError, match='candidate limit'):
        discover(Connector(), Provider(), replace(config, maximum_candidates=1), NOW)


def test_six_sheet_report_top_twenty_and_quality(product, listing, config):
    row = screened(product, listing, config)
    rows = [dict(row, product=f'Widget {i}', ranking_score=i) for i in range(25)]
    rows.append(dict(row, status='REJECT', quality_issues='Missing demand'))
    write_report(rows, config.output, discovery=True, job={'run_id': 17})
    wb = load_workbook(config.output)
    assert wb.sheetnames == ['Executive Summary', 'Top 20 Opportunities', 'All Candidates',
                            'Rejected Products', 'Data Quality and Missing Info', 'Job Execution Summary']
    top = wb['Top 20 Opportunities']
    assert top.max_row == 21 and top['B2'].value == 'Widget 24'
    assert top['L2'].hyperlink.target == product.url
    assert wb['Data Quality and Missing Info'].max_row == 27
    assert wb['Job Execution Summary']['B2'].value == 17
    assert 'MOCK' in wb['Executive Summary']['B1'].value
    assert 'Preparation' in [cell.value for cell in top[1]]
    wb.close()


def test_fixture_flow_and_stale_snapshots(config):
    rows = discover(CsvConnector(config.input), MockAmazonProvider(config.amazon_fixture), config, NOW)
    assert len(rows) == 3
    assert sum(r['status'] == 'REVIEW' for r in rows) == 1
    assert all(not r['purchase_ready'] for r in rows)
    stale = discover(CsvConnector(config.input), MockAmazonProvider(config.amazon_fixture), config, NOW+timedelta(days=2))
    assert not any(r['status'] == 'REVIEW' for r in stale)


def test_orchestrated_discovery_history(config):
    run_id = run_job(config)
    with sqlite3.connect(config.database) as db:
        assert db.execute('SELECT status,candidates FROM runs WHERE id=?', (run_id,)).fetchone() == ('succeeded', 3)
    wb = load_workbook(config.output)
    assert len(wb.sheetnames) == 6
    wb.close()


def test_bad_provider_fixture_fails_with_history(config, tmp_path):
    path = tmp_path/'bad.json'; path.write_text('not json')
    with pytest.raises(ValueError): run_job(replace(config, amazon_fixture=path))
    with sqlite3.connect(config.database) as db:
        assert db.execute('SELECT status FROM runs').fetchone() == ('failed',)

@pytest.mark.parametrize('change', [dict(mode='unknown'), dict(amazon_provider='live'),
    dict(inventory_budget=0), dict(maximum_data_age_hours=float('nan')), dict(top_count=0),
    dict(maximum_offers=-1), dict(maximum_candidates=0), dict(minimum_monthly_sales=1.5)])
def test_discovery_config_validation(config, change):
    with pytest.raises(ValueError): replace(config, **change)


def test_discovery_env_defaults(monkeypatch):
    monkeypatch.setenv('ARBITRAGE_MODE', 'discovery')
    monkeypatch.setenv('INVENTORY_BUDGET', '400')
    c = Config.from_env()
    assert c.input.name == 'discovery.csv'
    assert c.inventory_budget == 400
