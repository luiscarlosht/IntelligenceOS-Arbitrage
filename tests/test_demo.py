from contextlib import closing
from openpyxl import load_workbook

from app.demo import generate


def test_demo_screening_and_threshold_sensitivity(tmp_path, monkeypatch):
    # Ambient mail/provider settings must never affect this isolated workflow.
    monkeypatch.setenv('EMAIL_TO', 'must-not-send@example.org')
    monkeypatch.setenv('AMAZON_PROVIDER', 'live')
    monkeypatch.setattr('app.mailer.send_report', lambda *a: (_ for _ in ()).throw(AssertionError('mail')))
    monkeypatch.setattr('socket.socket.connect', lambda *a: (_ for _ in ()).throw(AssertionError('network')))
    monkeypatch.setattr('getpass.getpass', lambda *a: (_ for _ in ()).throw(AssertionError('credential')))
    result = generate(tmp_path)
    assert result['live_requests'] == 0
    scenarios = result['scenarios']
    assert set(scenarios['defaults']) == {'DEMO passing', 'DEMO boundary'}
    assert set(scenarios['budget_600']) - set(scenarios['defaults']) == {'DEMO over-budget'}
    assert set(scenarios['profit_4']) - set(scenarios['defaults']) == {'DEMO low-profit'}
    assert set(scenarios['roi_40']) - set(scenarios['defaults']) == {'DEMO low-roi'}
    rows = {r['product']: r for r in result['rows']}
    assert 'Package quantity' in rows['DEMO package-mismatch']['reason']
    assert rows['DEMO missing-fees']['profit'] is None
    assert 'Missing cost fulfillment_fee' in rows['DEMO missing-fees']['reason']
    assert 'Stale source' in rows['DEMO stale']['reason']
    assert all(not r['purchase_ready'] and r['data_mode'].startswith('DEMO') for r in rows.values())
    assert not result['authorization_verified']
    valid_issues = result['validation']['demo_authorized.csv']
    assert len(valid_issues) == 1 and 'Stale' in valid_issues[0]['issue']
    invalid = result['validation']['demo_invalid_authorized.csv']
    assert len(invalid) >= 8 and all(i['action'] for i in invalid)
    with closing(load_workbook(tmp_path / 'DEMO.xlsx')) as workbook:
        assert len(workbook.sheetnames) == 6
        assert all(name.startswith('DEMO') for name in workbook.sheetnames)
        assert workbook['DEMO Summary']['A1'].value.startswith('DEMO')
        assert workbook['DEMO Review Candidates'].max_row == 3
    assert not (tmp_path / 'unused.sqlite3').exists()
