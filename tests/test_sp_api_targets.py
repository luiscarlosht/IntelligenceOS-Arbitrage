from copy import deepcopy
from datetime import datetime, timezone
import json
from pathlib import Path
import sqlite3
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from app.amazon_catalog import MARKETPLACE
from app.amazon_targets import TargetStore, acquisition_ceiling
from app.observations import canonical, fingerprint, normalized
from app.sp_catalog_targets import (aggregate_targets, audit_saved, identity_priority,
    parse_target_item, sales_ranks, select_targets)
from app.sp_api_targets import make_plan, prepare, task_handoff
from app.sp_target_execution import run_prepared

NOW = datetime(2026, 10, 31, tzinfo=timezone.utc)


def item(asin='B000000001'):
    return dict(asin=asin, summaries=[dict(marketplaceId=MARKETPLACE, brand='Sharpie',
        itemName='Fixture fine point marker', manufacturer='Fixture manufacturer', modelNumber='M1',
        size='12 count', color='Black', packageQuantity=1, partNumber='P1', itemClassification='BASE_PRODUCT')],
        identifiers=[dict(marketplaceId=MARKETPLACE, identifiers=[
            dict(identifierType='UPC', identifier='071641300729')])],
        relationships=[dict(marketplaceId=MARKETPLACE, relationships=[])],
        attributes=dict(brand=[dict(marketplace_id=MARKETPLACE, value='Sharpie')],
            number_of_items=[dict(marketplace_id=MARKETPLACE, value=12)],
            unit_count=[dict(marketplace_id=MARKETPLACE, value=12, type={'value': 'count'})]),
        salesRanks=[dict(marketplaceId=MARKETPLACE, classificationRanks=[
            dict(classificationId='123', title='Fixture writing supplies', rank=1234)],
            displayGroupRanks=[dict(websiteDisplayGroup='office_product_display_on_website', title='Office', rank=999)])])


def snapshot(asin='B000000001'):
    return parse_target_item(item(asin), observed_at=NOW, provenance='synthetic fixture', is_mock=False)


def saved_report(directory, *, count=1, mismatch=False, received=True):
    rows, specs, attempts = [], [], []
    for i in range(count):
        value = snapshot(f'B{i:09d}')
        source = {'retailer': 'Walmart', 'source_context': {'retailer_item_id': str(i)}}
        rows.append(dict(ordinal=i + 1, source=source,
            search=dict(selected_asin=value['asin'], observed_at=NOW.isoformat(),
                        candidates=[dict(asin=value['asin'], title=value['catalog']['title'], brand='Sharpie', rank=i+1)]),
            detail=dict(value={'catalog': value['catalog'], 'package_attributes': {
                'brand': ['Sharpie'], 'size': ['12 count'], 'number_of_items': ['12'], 'unit_count': ['12.0']}},
                is_mock=False, observed_at=NOW.isoformat(), provenance='synthetic saved catalog')))
        specs.append(dict(ordinal=i + 1, source_fingerprint='bad' if mismatch else fingerprint(source),
            search={'params': {'marketplaceIds': MARKETPLACE}}, detail={'params': {'marketplaceIds': MARKETPLACE}}))
        for kind in ('search', 'detail'):
            attempts.append(dict(ordinal=i+1, kind=kind, status='response_received' if received else 'reserved'))
    document = dict(is_mock=False, status='succeeded', manifest={'marketplace_id': MARKETPLACE, 'requests': specs},
                    rows=rows, attempts=attempts)
    path = directory / 'candidates-fixture.json'
    path.write_text(canonical(document))
    return path


def test_parser_extends_identity_and_preserves_category_rank_snapshots():
    value = snapshot()
    assert value['summary_fields']['brand'] == 'Sharpie'
    assert value['summary_fields']['partNumber'] == 'P1'
    assert value['package_attributes']['unit_count'][0]['type']['value'] == 'count'
    ranks = value['sales_ranks']
    assert {r['category_id'] for r in ranks} == {'123', 'office_product_display_on_website'}
    assert all(r['observed_at'] == NOW.isoformat() and r['source_updated_at'] is None for r in ranks)
    target = aggregate_targets([value], now=NOW)[0]
    assert target.exact_package_evidence['unit_count']['value']['type']['value'] == 'count'
    changed = deepcopy(value)
    changed['sales_ranks'][0]['rank'] = 1
    assert identity_priority(target) == identity_priority(aggregate_targets([changed], now=NOW)[0])
    assert acquisition_ceiling(target) is None


@pytest.mark.parametrize('rank', [0, -1, True, '12'])
def test_rank_requires_positive_integer_and_category(rank):
    value = item()
    value['salesRanks'][0]['classificationRanks'][0]['rank'] = rank
    with pytest.raises(ValueError):
        sales_ranks(value, NOW, 'fixture')


def test_marketplace_scope_no_rank_inference_and_sanitization():
    value = item()
    value['salesRanks'][0]['marketplaceId'] = 'OTHER'
    value['attributes']['size'] = [{'marketplace_id': 'OTHER', 'value': 'wrong size'}]
    value['summaries'][0]['itemName'] = '=TOKEN'
    parsed = parse_target_item(value, observed_at=NOW, provenance='fixture', secrets=('TOKEN',))
    assert parsed['sales_ranks'] == []
    assert 'size' not in parsed['package_attributes']
    assert parsed['catalog']['title'] == '=[REDACTED]'
    assert parsed['catalog']['monthly_sales'] is None
    value['summaries'][0]['marketplaceId'] = 'OTHER'
    with pytest.raises(Exception):
        parse_target_item(value, observed_at=NOW, provenance='fixture')


def test_dedup_conflicts_and_all_identifiers_retained():
    first = snapshot()
    second = deepcopy(first)
    second['provenance'] = 'second fixture'
    second['summary_fields']['size'] = '24 count'
    second['catalog']['identifiers'].append(dict(identifier_type='UPC', value='036000291452', normalized_gtin='00036000291452'))
    targets = aggregate_targets([first, first, second], now=NOW)
    assert len(targets) == 1
    target = targets[0]
    assert target.exact_package_evidence['size']['value'] is None
    assert target.exact_package_evidence['size']['status'] == 'conflicting'
    assert {i.value for i in target.identifiers} == {'071641300729', '036000291452'}
    assert 'Ambiguous GTIN identifiers' in target.historical_metrics['mapping_issues']
    assert 'size' in target.missing_fields


@pytest.mark.parametrize('mismatch,received', [(True, True), (False, False)])
def test_saved_evidence_requires_binding_and_success(tmp_path, mismatch, received):
    saved_report(tmp_path, mismatch=mismatch, received=received)
    snapshots, audit = audit_saved(tmp_path)
    assert snapshots == []
    assert audit['known_search_asins'] == 1


def test_audit_and_artifacts_are_offline_formula_safe_and_bounded(tmp_path, monkeypatch):
    evidence = tmp_path / 'evidence'
    evidence.mkdir()
    source = saved_report(evidence, count=22)
    original = source.read_bytes()
    # This file is never opened; the evidence allowlist excludes credentials.
    (evidence / 'credentials.json').write_text('MUST NOT READ')
    network = Mock(side_effect=AssertionError('No network/credentials'))
    monkeypatch.setattr('urllib.request.OpenerDirector.open', network)
    monkeypatch.setattr('app.amazon_pilot.load_credentials', network)
    output = tmp_path / 'output'
    result = prepare(evidence, output, now=NOW)
    assert result['selected_targets'] == 20
    assert result['browser_jobs'] == 10
    assert result['live_requests'] == 0
    jobs = json.loads((output / 'browser-handoff.json').read_text())['jobs']
    assert len({j['job_id'] for j in jobs}) == 10
    assert all(j['location'] == '76063' and not j['application_verified_exact'] for j in jobs)
    assert all(j['acquisition_price_ceiling_usd'] is None for j in jobs)
    plan = json.loads((output / 'catalog-plan.json').read_text())
    assert plan['limits']['total'] == 11 and plan['expected_new_requests'] == 0
    assert len(plan['requests']) == 10
    assert all(r['params']['pageSize'] == 2 and 'salesRanks' in r['params']['includedData'] for r in plan['requests'])
    assert plan['limits']['pricing'] == plan['limits']['fees'] == plan['limits']['restrictions'] == 0
    workbook = load_workbook(output / 'AMAZON-TARGETS.xlsx')
    assert workbook['Targets'].max_row == 21
    assert not any(c.data_type == 'f' for s in workbook for row in s for c in row)
    assert source.read_bytes() == original
    assert (output / 'targets.sqlite3').stat().st_mode & 0o777 == 0o600
    with sqlite3.connect(output / 'targets.sqlite3') as db:
        assert db.execute('SELECT count(*) FROM amazon_search_jobs').fetchone()[0] == 10
    # Full saved seed means the future runner also returns before credentials or transport.
    result = run_prepared(output / 'catalog-plan.json', result['manifest_sha256'], 11,
                          private_path=evidence, credential_loader=network, transport=network)
    assert result['credential_files_read'] == 0 and result['live_requests'] == 0
    network.assert_not_called()
    with pytest.raises(FileExistsError):
        prepare(evidence, output, now=NOW)


def plan_file(tmp_path):
    evidence = tmp_path / 'private'
    evidence.mkdir(mode=0o700)
    _, audit = audit_saved(evidence)
    path = tmp_path / 'plan.json'
    path.write_text(canonical(make_plan(audit, [], evidence_directory=evidence)))
    import hashlib
    return evidence, path, hashlib.sha256(path.read_bytes()).hexdigest()


def test_future_runner_bounds_and_permanent_claims_with_mock_transport(tmp_path):
    evidence, path, digest = plan_file(tmp_path)
    calls = []
    def transport(request):
        calls.append(request.full_url)
        if 'api.amazon.com/auth' in request.full_url:
            return {'access_token': 'TEST_ONLY', 'token_type': 'bearer'}
        return {'numberOfResults': 0, 'items': [], 'pagination': {}}
    loader = Mock(return_value={'client_id': 'test', 'client_secret': 'test', 'refresh_token': 'test'})
    result = run_prepared(path, digest, 11, private_path=evidence, credential_loader=loader,
                          transport=transport, sleep=lambda _: None)
    assert len(calls) == 11 and result['attempts'] == 11 and result['live_requests'] == 0
    with sqlite3.connect(evidence / 'candidate-executions.sqlite3') as db:
        assert db.execute('SELECT count(*) FROM candidate_operations').fetchone()[0] == 11
    # Repeating/renaming the plan cannot release any consumed query claim.
    with pytest.raises(ValueError):
        run_prepared(path, digest, 11, private_path=evidence, credential_loader=loader, transport=transport)
    assert loader.call_count == 1


def test_quota_stop_is_sanitized_and_never_retried(tmp_path):
    from app.amazon_diagnostics import AmazonFailure
    evidence, path, digest = plan_file(tmp_path)
    calls = []
    def transport(request):
        calls.append(request.full_url)
        if 'api.amazon.com/auth' in request.full_url:
            return {'access_token': 'TEST_ONLY', 'token_type': 'bearer'}
        raise AmazonFailure('SECRET DO NOT LOG', reason='http_error', http_status=429, codes=('QuotaExceeded',))
    with pytest.raises(AmazonFailure):
        run_prepared(path, digest, 11, private_path=evidence, transport=transport,
            credential_loader=lambda _: {'client_id': 'test', 'client_secret': 'test', 'refresh_token': 'test'})
    assert len(calls) == 2
    report = next(evidence.glob('amazon-first-catalog-*.json')).read_text()
    assert 'SECRET DO NOT LOG' not in report and 'TEST_ONLY' not in report
    with sqlite3.connect(evidence / 'candidate-executions.sqlite3') as db:
        assert db.execute('SELECT operation FROM amazon_cooldowns').fetchone()[0] == 'searchCatalogItems'


def test_wrong_manifest_hash_never_loads_credentials(tmp_path):
    evidence, path, _ = plan_file(tmp_path)
    loader = Mock(side_effect=AssertionError('No credentials'))
    with pytest.raises(ValueError):
        run_prepared(path, 'bad', 11, private_path=evidence, credential_loader=loader)
    loader.assert_not_called()


def test_refresh_requires_new_version_and_reason(tmp_path):
    targets = select_targets(aggregate_targets([snapshot()], now=NOW))
    job = task_handoff(targets)[0]
    store = TargetStore(tmp_path / 'pilot.sqlite3')
    store.save_target(targets[0])
    assert store.save_job(job)
    assert not store.save_job(job)
    changed = dict(job, query_version='v2')
    changed['job_id'] = 'search:' + fingerprint({k: changed[k] for k in ('target_id', 'retailer', 'location', 'query_version')})
    with pytest.raises(ValueError):
        store.save_job(changed)
    assert store.save_job(changed, prior_job_id=job['job_id'], refresh_reason='new model evidence')
    for outcome in ('no_result', 'mismatch', 'unavailable', 'unresolved', 'matched'):
        store.record_outcome(job['job_id'], outcome=outcome, observed_at=NOW, query='M1',
            evidence_reference='fixture:' + outcome, elapsed_seconds=1, cost_usd=None)
    metrics = store.metrics()
    assert metrics['query_attempts'] == 5 and metrics['elapsed_seconds'] == 5
    assert metrics['exact_match_jobs'] == metrics['profitable_match_jobs'] == 0
    assert metrics['total_cost_usd'] is None


def test_independent_verification_cannot_bypass_variant_brand_or_costs(tmp_path):
    from dataclasses import replace
    from decimal import Decimal
    from app.amazon_catalog import parse_catalog
    from app.amazon_targets import verify_observation
    from app.config import Config
    from app.models import RetailProduct
    from app.profitability import COST_FIELDS
    raw = item()
    raw['summaries'][0]['style'] = 'Fine'
    target = aggregate_targets([parse_target_item(raw, observed_at=NOW, provenance='fixture')], now=NOW)[0]
    listing = parse_catalog(raw, expected_asin=target.asin, is_mock=False).value
    listing = replace(listing, seller_eligible=True, sale_price=Decimal('30'),
        costs={key: Decimal('0') for key in COST_FIELDS}, verified_costs=COST_FIELDS,
        monthly_sales=50, offer_count=3, retrieved_at=NOW)
    product = RetailProduct('Walmart', 'fixture', 'https://www.walmart.com/ip/fixture/123', listing.gtin,
        listing.manufacturer, listing.model, 1, Decimal('10'), 'in_stock', NOW, 'fixture',
        asin_hint=target.asin, variation=listing.variation,
        source_context={'brand': 'Sharpie', 'observed_zip': '76063'})
    config = Config(tmp_path/'input.csv', tmp_path/'output.xlsx', tmp_path/'history.sqlite3')
    job = task_handoff([target])[0]
    assert verify_observation(target, job, product, listing, config, NOW)['verified_profitable']
    changed_variant = canonical({'color': 'Black', 'size': '12 count', 'style': 'UltraFine'})
    assert not verify_observation(target, job, replace(product, variation=changed_variant),
        replace(listing, variation=changed_variant), config, NOW)['application_verified_exact']
    assert not verify_observation(target, job, product, replace(listing, costs={}), config, NOW)['verified_profitable']
    assert not verify_observation(target, job, replace(product, source_context={}), listing, config, NOW)['application_verified_exact']
