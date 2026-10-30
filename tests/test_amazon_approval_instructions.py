from openpyxl import Workbook, load_workbook
from app.amazon_details import parse_restrictions
from app.provider_contracts import SellerContext
from app.amazon_catalog import MARKETPLACE
from app.amazon_approval_report import write_approval_links


def test_preserves_reason_messages_and_links_as_literal_cells(tmp_path):
    links = [{'resource': 'https://sellercentral.amazon.com/approval?a=1&b=2', 'title': '=literal', 'verb': 'GET'},
             {'resource': 'javascript:alert(1)', 'title': 'unsafe'},
             {'resource': 'https://user:password@example.com', 'title': 'unsafe'}]
    data = parse_restrictions({'restrictions': [{'marketplaceId': MARKETPLACE, 'reasons': [
        {'reasonCode': 'APPROVAL_REQUIRED', 'message': '=Read instructions', 'links': links},
        {'reasonCode': 'FUTURE_REASON', 'message': 'Keep this'}]}]},
        SellerContext('FAKE', MARKETPLACE, 'new_new'), is_mock=True)
    assert data['reasons'][0]['links'] == links
    assert data['reasons'][1]['reasonCode'] == 'FUTURE_REASON'
    assert data['approval_required'] is True
    assert data['listing_permission']['value'] is False
    assert data['fulfillment_eligibility'] is None
    wb = Workbook()
    write_approval_links(wb, [{'candidate_asin': 'B000000001', 'approval_messages': '=Read instructions',
                             'approval_links': links}])
    path = tmp_path/'test.xlsx'; wb.save(path)
    sheet = load_workbook(path)['Amazon approval instructions']
    assert sheet['C2'].value == '=Read instructions' and sheet['C2'].data_type == 's'
    assert sheet['D2'].data_type == 's'
    assert sheet['E2'].hyperlink.target == links[0]['resource']
    assert sheet['E3'].hyperlink is None and sheet['E4'].hyperlink is None


def test_missing_instructions_stay_unknown():
    data = parse_restrictions({'restrictions': []}, SellerContext('FAKE', MARKETPLACE, 'new_new'), is_mock=True)
    assert data['approval_instructions'] is None
    assert data['approval_required'] is False
    assert data['fulfillment_eligibility'] is None
