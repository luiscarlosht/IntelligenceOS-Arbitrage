import json
import logging
import sqlite3
import threading
from dataclasses import replace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from app.config import Config
from app.pipeline import evaluate, load_candidates, ROOT
from app.reporting import write_report
from app.orchestrator import run_job, serve, JsonFormatter
from app.mailer import send_report

@pytest.fixture
def row():
    return dict(source='sample', product='widget', upc='000123', source_url='https://example.org',
                amazon_asin='A123', destination='Amazon', buy_cost='8', sale_price='25',
                referral_fee='3.75', fulfillment_fee='4.5', inbound_shipping='1.25',
                other_cost='0.5', eligible='yes', exact_match='yes')

@pytest.fixture
def config(tmp_path):
    return Config(ROOT / 'sample_data/candidates.csv', tmp_path / 'report.xlsx', tmp_path / 'history.sqlite3')

@pytest.mark.parametrize('value', ['nan', 'inf', '-inf', '-1', 'bad', '0'])
def test_invalid_buy_cost(row, value):
    row['buy_cost'] = value
    assert evaluate(row)['status'] == 'ERROR'

def test_roi_and_threshold_boundaries(row):
    result = evaluate(row)
    assert result['roi'] == 71.8
    assert evaluate(row, minimum_roi=72)['status'] == 'REJECT'
    assert evaluate(row, minimum_profit=7)['status'] == 'REVIEW'
    assert evaluate(row, minimum_profit=7.01)['status'] == 'REJECT'
    row.update(buy_cost='10', sale_price='15', referral_fee='0', fulfillment_fee='0', inbound_shipping='0', other_cost='0')
    assert evaluate(row)['status'] == 'REVIEW'  # exactly $5 and 50%

def test_match_and_missing_fields(row):
    row['exact_match'] = 'no'
    assert evaluate(row)['status'] == 'REJECT'
    del row['sale_price']
    assert evaluate(row)['status'] == 'ERROR'

def test_missing_csv_headers(tmp_path):
    path = tmp_path / 'bad.csv'
    path.write_text('product\nwidget\n')
    with pytest.raises(ValueError, match='Missing CSV columns'):
        load_candidates(path)

def test_report_contents_and_formula_safety(row, tmp_path):
    row['product'] = '=HYPERLINK("https://example.org")'
    path = tmp_path / 'report.xlsx'
    write_report([evaluate(row)], path)
    wb = load_workbook(path)
    assert wb.sheetnames == ['Executive Summary', 'Top Opportunities', 'All Candidates', 'Rejected Products']
    sheet = wb['All Candidates']
    assert sheet['B2'].data_type == 's'
    assert sheet['C2'].value == '000123'
    assert sheet['H2'].value == 7
    assert sheet['I2'].value == 71.8
    assert sheet.freeze_panes == 'A2'
    assert wb['Executive Summary']['B3'].value == 1
    wb.close()

def test_empty_report(tmp_path):
    path = tmp_path / 'empty.xlsx'
    write_report([], path)
    wb = load_workbook(path)
    assert wb['All Candidates'].max_row == 1
    wb.close()

def records(config):
    with sqlite3.connect(config.database) as db:
        return db.execute('SELECT status,candidates,review_candidates,report_path,error_type FROM runs ORDER BY id').fetchall()

def test_jobs_persist_and_repeat(config):
    assert run_job(config) == 1
    assert run_job(config) == 2
    assert [r[:3] for r in records(config)] == [('succeeded', 3, 1)] * 2
    wb = load_workbook(config.output)
    assert wb['All Candidates'].max_row == 4
    wb.close()

def test_failed_job_history(config):
    config = replace(config, input=config.input.parent / 'missing.csv')
    with pytest.raises(FileNotFoundError): run_job(config)
    assert records(config)[0][0] == 'failed'
    assert records(config)[0][4] == 'FileNotFoundError'

def test_email_failure_preserves_report_and_history(config, monkeypatch):
    def fail(*args): raise OSError('redacted')
    monkeypatch.setattr('app.orchestrator.send_report', fail)
    with pytest.raises(OSError): run_job(config)
    assert config.output.exists()
    assert records(config)[0][0] == 'failed'
    assert records(config)[0][3] == str(config.output)

@pytest.mark.parametrize('security', ['starttls', 'ssl'])
def test_email(config, monkeypatch, security):
    config = replace(config, smtp_host='smtp.example.org', email_from='from@example.org',
                     email_to=('to@example.org',), smtp_user='user', smtp_password='secret', smtp_security=security)
    config.output.write_bytes(b'report')
    factory = MagicMock()
    monkeypatch.setattr('app.mailer.smtplib.SMTP_SSL' if security == 'ssl' else 'app.mailer.smtplib.SMTP', factory)
    assert send_report(config, config.output, 1)
    client = factory.return_value.__enter__.return_value
    assert client.starttls.call_count == (1 if security == 'starttls' else 0)
    client.login.assert_called_once_with('user', 'secret')
    message = client.send_message.call_args.args[0]
    assert list(message.iter_attachments())[0].get_filename() == 'report.xlsx'

def test_disabled_email(config):
    assert send_report(config, config.output, 1) is False

@pytest.mark.parametrize('change', [dict(interval=-1), dict(minimum_roi=float('nan')),
    dict(smtp_security='plain'), dict(email_to=('a@b',)), dict(smtp_user='user'), dict(smtp_port=0)])
def test_bad_config(config, change):
    with pytest.raises(ValueError): replace(config, **change)

def test_environment_config(monkeypatch):
    monkeypatch.setenv('MINIMUM_ROI', '60')
    monkeypatch.setenv('JOB_INTERVAL_SECONDS', '120')
    monkeypatch.setenv('EMAIL_TO', '')
    assert Config.from_env().minimum_roi == 60
    assert Config.from_env().interval == 120

def test_scheduler_continues_after_failure(config, monkeypatch):
    stop = threading.Event()
    calls = []
    def job(_):
        calls.append(1)
        if len(calls) == 1: raise OSError()
        stop.set()
    monkeypatch.setattr('app.orchestrator.run_job', job)
    serve(replace(config, interval=0.001), stop)
    assert len(calls) == 2

def test_one_shot_failure_propagates(config, monkeypatch):
    def fail(_): raise OSError()
    monkeypatch.setattr('app.orchestrator.run_job', fail)
    with pytest.raises(OSError): serve(config, threading.Event())

def test_json_log():
    record = logging.LogRecord('arbitrage', logging.INFO, '', 0, 'job_started', (), None)
    record.run_id = 4
    result = json.loads(JsonFormatter().format(record))
    assert result['event'] == 'job_started'
    assert result['run_id'] == 4
