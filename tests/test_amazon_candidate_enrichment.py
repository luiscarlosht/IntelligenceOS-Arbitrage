"""Synthetic responses only. No network or real credentials used."""
from copy import deepcopy
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from app import amazon_candidate_enrichment as runner
from app import amazon_pilot as pilot
from app.amazon_catalog import MARKETPLACE
from app.amazon_details import parse_fees
from app.candidate_review_report import build_review, regenerate
from app.observations import normalized
from app.provider_contracts import SellerContext
from app.provider_runtime import ProviderFailure
from app.reporting import write_report

CREDS = dict(client_id='mock-client', client_secret='mock-secret', refresh_token='mock-refresh')
SELLER = 'MOCKSELLER'
STAMP = '2026-10-28T12:00:00+00:00'


def fixture(tmp_path, *, mocked=True):
    rows=[]
    for i in range(1,11):
        # Two source listings select one candidate; all ten rows must survive.
        asin=f'B{(4 if i==7 else i):09d}'
        source=dict(retailer='eBay',name=f'Fixture utensil {i}',url=f'https://www.ebay.com/itm/{100000000000+i}',
                    source_context=dict(search_price='10', source_shipping_cost=None if i==1 else '0'),
                    retrieved_at=STAMP,provenance='Mock source export',gtin='',manufacturer='',model='')
        rows.append(dict(ordinal=i,source=source,status='ambiguous',exact=False,
            search=dict(selected_asin=asin,candidates=[dict(asin=asin,title='Synthetic candidate')],complete=False),
            detail=dict(value=dict(catalog=dict(asin=asin,is_mock=mocked,title='Synthetic candidate')),
                        is_mock=mocked,observed_at=STAMP)))
    path=tmp_path/'candidates.json'
    path.write_text(json.dumps(dict(is_mock=mocked,status='succeeded',rows=rows,live_requests=21,attempts=[{}]*21)))
    return path


def pricing(asin, *, shipping='1', landed='21'):
    price=lambda value: dict(CurrencyCode='USD',Amount=float(value)) if value is not None else None
    return dict(payload=dict(status='Success',ASIN=asin,marketplaceId=MARKETPLACE,ItemCondition='New',
        Identifier=dict(ASIN=asin,MarketplaceId=MARKETPLACE,ItemCondition='New'),Offers=[{}],
        Summary=dict(TotalOfferCount=1,LowestPrices=[dict(condition='New',fulfillmentChannel='Amazon',
            ListingPrice=price('20'),Shipping=price(shipping),LandedPrice=price(landed))])))


def fees(asin, request):
    spec=request['FeesEstimateRequest']
    return dict(payload=dict(FeesEstimateResult=dict(Status='Success',
        FeesEstimateIdentifier=dict(MarketplaceId=MARKETPLACE,SellerId=SELLER,IdType='ASIN',IdValue=asin,
            IsAmazonFulfilled=spec['IsAmazonFulfilled'],SellerInputIdentifier=spec['Identifier'],
            PriceToEstimateFees=spec['PriceToEstimateFees']),
        FeesEstimate=dict(TimeOfFeesEstimation=STAMP,TotalFeesEstimate=dict(CurrencyCode='USD',Amount=8),
            FeeDetailList=[dict(FeeType='ReferralFee',FinalFee=dict(CurrencyCode='USD',Amount=3)),
                           dict(FeeType='FBAFees',FinalFee=dict(CurrencyCode='USD',Amount=5))]))))


class Transport:
    def __init__(self, failure=None, shipping='1', reason=None):
        self.calls=[];self.failure=failure;self.shipping=shipping;self.reason=reason

    def __call__(self, request):
        self.calls.append(request)
        url=urlsplit(request.full_url);params=parse_qs(url.query)
        if url.hostname=='api.amazon.com':
            if self.failure=='authentication':raise ProviderFailure('mock-secret')
            return dict(access_token='mock-access',token_type='bearer')
        if url.path.endswith('/offers'):
            if self.failure=='pricing':raise ProviderFailure('mock-secret')
            return pricing(url.path.split('/')[-2],shipping=self.shipping)
        if url.path.endswith('/restrictions'):
            assert params['sellerId']==[SELLER]
            if self.failure=='restrictions':raise ProviderFailure('mock-secret')
            return dict(restrictions=[dict(marketplaceId=MARKETPLACE,conditionType='new_new',
                reasons=[dict(reasonCode=self.reason,message='mock-secret')])] if self.reason else [])
        if url.path.endswith('/feesEstimate'):
            if self.failure=='fees':raise ProviderFailure('mock-secret')
            body=json.loads(request.data)
            assert body['FeesEstimateRequest']['PriceToEstimateFees']['Shipping']['Amount']==0
            return fees(url.path.split('/')[-2],body)
        pytest.fail('Unplanned request')


def run(tmp_path, *, failure=None, fixed=None, shipping='1', reason=None):
    saved=fixture(tmp_path);manifest=runner.plan(saved,fixed_fee_prices=fixed)
    transport=Transport(failure,shipping,reason)
    clock = [0.]
    def advance(seconds): clock[0] += seconds
    result=runner.execute(saved,manifest,tmp_path/'private',transport=transport,
        credential_loader=lambda _:CREDS,seller_prompt=lambda:SELLER,is_mock=True,sleep=advance,monotonic=lambda:clock[0])
    return saved,manifest,transport,result


def test_dedup_manifest_requests_and_private_preserved_evidence(tmp_path):
    saved,manifest,transport,report=run(tmp_path)
    assert manifest['budget']==dict(authentication=1,pricing=9,fees=9,restrictions=9,maximum_requests=28,retries=0)
    assert len(transport.calls)==28 and report['status']=='succeeded'
    assert next(r for r in manifest['requests'] if r['asin']=='B000000004')['source_ordinals']==[4,7]
    path=next((tmp_path/'private').glob('candidate-enrichment-*.json'))
    raw=path.read_text();before=path.read_bytes()
    assert path.stat().st_mode & 0o777==0o600
    assert all(secret not in raw for secret in (*CREDS.values(),'mock-access',SELLER))
    assert json.loads(raw)==report and report['live_requests']==0
    with pytest.raises(FileExistsError):
        runner.execute(saved,manifest,tmp_path/'private',transport=lambda _:pytest.fail('network'),
            credential_loader=lambda _:pytest.fail('credentials'),seller_prompt=lambda:pytest.fail('prompt'),is_mock=True)
    assert path.read_bytes()==before


@pytest.mark.parametrize('failure',['pricing','fees','restrictions','authentication'])
def test_independent_stage_failures_no_retries(tmp_path,failure):
    saved,manifest,transport,report=run(tmp_path,failure=failure,fixed={'B000000001':'19.99'})
    first=report['rows'][0]
    if failure=='authentication':
        assert len(transport.calls)==1 and report['status']=='stopped_no_retry'
        assert all(first[k]['status']=='not_attempted' for k in ('pricing','fees','restrictions'))
    else:
        assert report['status']=='partial' and first[failure]['status']=='failed'
        assert first[failure]['data'] is None
        assert all(first[k]['status']=='succeeded' for k in ('pricing','fees','restrictions') if k!=failure)
        assert len(transport.calls)==(20 if failure=='pricing' else 28)
        if failure=='pricing':
            assert report['rows'][1]['fees']['status']=='skipped_missing_price'
    assert len(report['attempts'])==len(transport.calls)
    assert 'mock-secret' not in json.dumps(report)


def test_report_columns_zero_and_unknown_costs_and_no_purchase_approval(tmp_path):
    saved,manifest,transport,report=run(tmp_path,shipping=None,reason='APPROVAL_REQUIRED')
    original=json.loads(saved.read_text())
    rows,review,evidence=build_review(original,amazon_enrichment=report)
    first=review['comparisons'][0]
    assert first['amazon_item_price']==20 and first['amazon_shipping'] is None
    # A reported landed price must not be used to invent missing shipping.
    assert first['amazon_landed_price']==21
    assert first['fee_sale_price']==20 and first['fee_buyer_shipping']==0
    assert first['estimated_referral_fee']==3 and first['estimated_fba_fee']==5 and first['estimated_total_fees']==8
    assert first['eligibility_marketplace']==MARKETPLACE and first['eligibility_condition']=='new_new'
    assert 'APPROVAL_REQUIRED' in first['approval_requirements']
    assert first['inbound_fba_shipping'] is first['merchant_fulfillment_cost'] is None
    assert all(row['status']=='REJECT' and row['profit'] is None and row['purchase_ready'] is False for row in rows)
    assert len(rows)==10 and review['Candidate-only rows']==10
    assert review['comparisons'][3]['candidate_asin']==review['comparisons'][6]['candidate_asin']
    output=tmp_path/'mock.xlsx'
    write_report(rows,output,discovery=True,candidate_review=review,enrichment_rows=evidence)
    wb=load_workbook(output);sheet=wb['Candidate Comparison']
    headings={c.value:c.column for c in sheet[1]}
    assert sheet.max_row==11
    assert sheet['E2'].hyperlink
    assert wb['Amazon approval instructions']['C2'].value == 'mock-secret'
    assert sheet.cell(2,headings['Amazon buyer-paid shipping (USD)']).value is None
    assert sheet.cell(2,headings['Fee-estimate buyer-paid shipping (USD)']).value==0
    assert sheet.cell(2,headings['Inbound-to-FBA shipping cost (USD)']).value is None
    assert not any(c.data_type=='f' for s in wb for r in s for c in r)
    wb.close()


@pytest.mark.parametrize('change',['asin','marketplace','condition','price','shipping','fulfillment','mock','restriction_scope'])
def test_incompatible_saved_evidence_rejected(tmp_path,change):
    saved,_,_,report=run(tmp_path)
    item=report['rows'][0];quote=item['fees']['data']['quote']['value']
    if change=='asin':quote['asin']='B099999999'
    if change=='marketplace':quote['context']['marketplace_id']='OTHER'
    if change=='condition':quote['context']['condition']='used_like_new'
    if change=='price':quote['sale_price']='37.25'
    if change=='shipping':item['fees']['data']['buyer_shipping_scenario']='1'
    if change=='fulfillment':quote['fulfillment']='FBM'
    if change=='mock':item['pricing']['data']['is_mock']=False
    if change=='restriction_scope':item['condition']='used_like_new'
    with pytest.raises(ValueError):build_review(json.loads(saved.read_text()),amazon_enrichment=report)


def test_fee_parser_supports_candidate_asin_and_explicit_scenario(tmp_path):
    asin='B00OTRLKAE';context=SellerContext(SELLER,MARKETPLACE)
    request=dict(FeesEstimateRequest=dict(IsAmazonFulfilled=True,Identifier='candidate-fee',
        PriceToEstimateFees=dict(ListingPrice=dict(CurrencyCode='USD',Amount=19.99),Shipping=dict(CurrencyCode='USD',Amount=2))))
    document=fees(asin,request)
    result=parse_fees(document,context,Decimal('19.99'),'candidate-fee',is_mock=True,
                      expected_asin=asin,buyer_shipping=Decimal('2'))
    assert result['quote']['value']['asin']==asin and result['buyer_shipping_scenario']==2
    with pytest.raises(ProviderFailure):
        parse_fees(document,context,Decimal('37.25'),'candidate-fee',is_mock=True,expected_asin=asin,buyer_shipping=Decimal('2'))
    document['payload']['FeesEstimateResult']['FeesEstimate']['FeeDetailList']=[]
    result=parse_fees(document,context,Decimal('19.99'),'candidate-fee',is_mock=True,expected_asin=asin,buyer_shipping=Decimal('2'))
    assert result['quote']['value']['referral_fee'] is result['quote']['value']['fulfillment_fee'] is None
    assert result['total_fees_estimate']==8


def test_manifest_mutation_before_prompt_credentials_or_network(tmp_path):
    saved=fixture(tmp_path);manifest=runner.plan(saved)
    manifest['requests'][0]['fees']['sale_price']='37.25'
    with pytest.raises(ValueError,match='Manifest differs'):
        runner.execute(saved,manifest,tmp_path/'private',transport=lambda _:pytest.fail('network'),
            credential_loader=lambda _:pytest.fail('credentials'),seller_prompt=lambda:pytest.fail('prompt'),is_mock=True)
    assert not (tmp_path/'private').exists()


def test_interrupt_saves_prior_stage_and_prevents_retry(tmp_path):
    saved=fixture(tmp_path);manifest=runner.plan(saved);transport=Transport()
    def interrupted(req):
        if len(transport.calls)==2:raise KeyboardInterrupt('mock-secret')
        return transport(req)
    report=runner.execute(saved,manifest,tmp_path/'private',transport=interrupted,
        credential_loader=lambda _:CREDS,seller_prompt=lambda:SELLER,is_mock=True,sleep=lambda _:None)
    assert report['status']=='stopped_no_retry'
    assert report['rows'][0]['pricing']['status']=='succeeded'
    assert report['rows'][0]['restrictions']['status']=='interrupted'
    assert len(report['attempts'])==3
    assert 'mock-secret' not in json.dumps(report)


def test_hidden_prompt_failure_before_credential_load_and_alarm(tmp_path):
    saved=fixture(tmp_path);manifest=runner.plan(saved)
    def fail_prompt():raise ValueError('mock-secret')
    with patch.object(runner.signal,'alarm') as alarm:
        report=runner.execute(saved,manifest,tmp_path/'private',transport=lambda _:pytest.fail('network'),
            credential_loader=lambda _:pytest.fail('credentials'),seller_prompt=fail_prompt,is_mock=True)
    assert not report['attempts'] and report['status']=='stopped_no_retry'
    alarm.assert_not_called()


def test_scoped_no_restrictions_remains_candidate_only(tmp_path):
    saved,_,_,report=run(tmp_path)
    rows,review,_=build_review(json.loads(saved.read_text()),amazon_enrichment=report)
    assert 'UNRESTRICTED' in review['comparisons'][0]['restriction_status']
    assert all(row['matching_confidence']==0 and row['profit'] is None for row in rows)
    assert review['comparisons'][0]['overall_fba_eligibility']=='UNKNOWN'
    assert 'not purchase approval' in review['comparisons'][0]['analysis_notice']


def test_elapsed_deadline_caps_requests_and_restores_alarm(tmp_path):
    saved=fixture(tmp_path);manifest=runner.plan(saved)
    clock=iter([0,0,300])
    transport=Transport()
    with patch.object(runner.signal,'signal',return_value=runner.signal.SIG_DFL) as handler, \
         patch.object(runner.signal,'alarm') as alarm:
        report=runner.execute(saved,manifest,tmp_path/'private',transport=transport,
            credential_loader=lambda _:CREDS,seller_prompt=lambda:SELLER,is_mock=True,
            sleep=lambda _:None,monotonic=lambda:next(clock))
    assert report['status']=='stopped_no_retry' and not transport.calls
    assert alarm.call_args_list[0].args==(300,) and alarm.call_args_list[-1].args==(0,)
    assert handler.call_args_list[-1].args==(runner.signal.SIGALRM,runner.signal.SIG_DFL)


def test_offline_plan_never_prompts_loads_credentials_or_calls_http(tmp_path):
    saved=fixture(tmp_path)
    with patch.object(pilot,'load_credentials',side_effect=AssertionError('credentials')), \
         patch.object(pilot,'request_json',side_effect=AssertionError('network')), \
         patch.object(runner,'prompt_seller_id',side_effect=AssertionError('prompt')):
        manifest=runner.plan(saved)
    assert manifest['live_requests']==0 and not manifest['execution_authorized']


def test_regenerate_with_saved_scoped_enrichment_preserves_all_inputs(tmp_path):
    # Synthetic live-shaped provenance exercises the production regeneration guard.
    saved=fixture(tmp_path,mocked=False)
    original=json.loads(saved.read_text())
    items=[dict(name=r['source']['name'],url=r['source']['url'],shop_name=None,price='10',
                shipping_costs=r['source']['source_context']['source_shipping_cost']) for r in original['rows']]
    export=tmp_path/'source-export.json'
    export.write_text(json.dumps(dict(status='finished',results=[dict(success=True,
        query=dict(source='ebay',country='us',topic='search_results'),metadata=dict(page_count_live=1),
        content=dict(search_results=items))])))
    digest=hashlib.sha256(export.read_bytes()).hexdigest()
    for row in original['rows']:row['source']['provenance']='Synthetic saved export sha256='+digest
    saved.write_text(json.dumps(original))
    manifest=runner.plan(saved);transport=Transport()
    report=runner.execute(saved,manifest,tmp_path/'private',transport=transport,
        credential_loader=lambda _:CREDS,seller_prompt=lambda:SELLER,is_mock=False,sleep=lambda _:None)
    enrichment=next((tmp_path/'private').glob('candidate-enrichment-*.json'))
    old=tmp_path/'original.xlsx';old.write_bytes(b'original workbook sentinel')
    before={p:p.read_bytes() for p in (saved,export,enrichment,old)}
    output=tmp_path/'enriched.xlsx'
    result=regenerate(saved,export,old,output,amazon_enrichment=enrichment)
    assert result['rows']==10 and result['new_api_calls']==0
    assert all(p.read_bytes()==value for p,value in before.items())
    wb=load_workbook(output);sheet=wb['Candidate Comparison'];head={c.value:c.column for c in sheet[1]}
    assert sheet.cell(2,head['Amazon observed item price (USD)']).value==20
    assert sheet.cell(2,head['Amazon buyer-paid shipping (USD)']).value==1
    assert sheet.cell(2,head['Estimated referral fee (USD)']).value==3
    assert all(sheet.cell(r,head['Overall seller / FBA eligibility']).value=='UNKNOWN' for r in range(2,12))
    assert sheet['F5'].value==sheet['F8'].value
    wb.close()
    report['manifest']['saved_report_sha256']='0'*64
    enrichment.write_text(json.dumps(report))
    with pytest.raises(ValueError,match='manifest'):
        regenerate(saved,export,old,tmp_path/'bad.xlsx',amazon_enrichment=enrichment)
    assert not (tmp_path/'bad.xlsx').exists()


def test_saved_price_scope_and_zero_shipping(tmp_path):
    saved,_,_,report=run(tmp_path,shipping='0')
    original=json.loads(saved.read_text())
    _,review,_=build_review(original,amazon_enrichment=report)
    assert review['comparisons'][0]['amazon_shipping']==0
    for field,value in [('asin','B099999999'),('marketplace_id','OTHER'),('condition','Used'),('customer_type','Business')]:
        changed=deepcopy(report);changed['rows'][0]['pricing']['data']['value'][field]=value
        with pytest.raises(ValueError,match='Pricing evidence scope'):
            build_review(original,amazon_enrichment=changed)


def test_throttle_preserves_attempts_and_independent_stages(tmp_path):
    from app.amazon_diagnostics import AmazonFailure
    saved=fixture(tmp_path); manifest=runner.plan(saved)
    transport=Transport(); clock=[0.]
    def advance(seconds): clock[0] += seconds
    def throttled(request):
        if request.full_url.split('?')[0].endswith('/offers'):
            error=AmazonFailure('private', reason='http_error', http_status=429, codes=('QuotaExceeded',))
            error.rate_headers={'Retry-After':'1000'}
            raise error
        return transport(request)
    report=runner.execute(saved,manifest,tmp_path/'private',transport=throttled,
        credential_loader=lambda _:CREDS,seller_prompt=lambda:SELLER,is_mock=True,
        sleep=advance,monotonic=lambda:clock[0])
    assert report['rows'][0]['pricing']['failure']['http_status']==429
    assert report['rows'][0]['restrictions']['status']=='succeeded'
    assert report['rows'][1]['pricing']['failure']['reason']=='pacing_deferred'
    assert all(row['restrictions']['status']=='succeeded' for row in report['rows'])
    assert len(report['attempts']) >= 1+2*len(report['rows'])
    assert report['attempts'][1]['status']=='failed'
    assert report['attempts'][1]['failure']['pacing']['not_before_monotonic']==1000


def test_returned_approval_links_map_to_correct_candidate_report(tmp_path):
    saved, manifest, transport, report = run(tmp_path, reason='APPROVAL_REQUIRED')
    first = report['rows'][0]
    first['restrictions']['data']['reasons'][0]['links'] = [
        dict(resource='https://sellercentral.amazon.com/approval?asin='+first['asin'],
             title='=Amazon instructions', verb='GET')]
    _, review, _ = build_review(json.loads(saved.read_text()), amazon_enrichment=report)
    comparison = next(r for r in review['comparisons'] if r['candidate_asin'] == first['asin'])
    assert comparison['approval_messages'] == 'mock-secret'
    assert comparison['approval_links'][0]['reason_code'] == 'APPROVAL_REQUIRED'
    assert comparison['approval_links'][0]['resource'].endswith(first['asin'])
    assert comparison['overall_fba_eligibility'] == 'UNKNOWN'
