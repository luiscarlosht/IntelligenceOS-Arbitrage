"""Fictional end-to-end evidence lineage; no accounts, credentials or HTTP."""
import csv
from dataclasses import replace
from datetime import timedelta, timezone
from decimal import Decimal
import json
from pathlib import Path
import sqlite3
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from app.config import ROOT
from app.connectors import CsvConnector
from app.marketplace import MockAmazonProvider
from app.observations import ObservationStore, canonical, fingerprint
from app import offline_pilot as pilot
from app.provider_contracts import Evidence


def run(tmp_path, source=None, **kwargs):
    return pilot.run_offline(source or ROOT/'sample_data/pilot_source_fixture.csv',
        ROOT/'sample_data/pilot_catalog_fixture.json', ROOT/'sample_data/pilot_keepa_fixture.json',
        tmp_path/'state', **kwargs)


def source_rows(tmp_path, count=1, duplicate=False, **changes):
    with (ROOT/'sample_data/pilot_source_fixture.csv').open() as stream:
        reader = csv.DictReader(stream)
        headers, rows = reader.fieldnames, list(reader)
    rows = [dict(rows[0], **changes) for _ in range(count)] if duplicate else [dict(row, **changes) for row in rows[:count]]
    target = tmp_path/'input.csv'
    with target.open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    return target


def test_pipeline_fixture_evidence_and_report(tmp_path, monkeypatch):
    network = Mock(side_effect=AssertionError('Network forbidden'))
    monkeypatch.setattr('urllib.request.OpenerDirector.open', network)
    monkeypatch.setattr('app.gmail_api.read_private', network)
    monkeypatch.setattr('app.mailer.send_report', network)
    monkeypatch.setenv('EMAIL_TO', 'must-not-send@example.com')
    monkeypatch.setenv('AMAZON_PROVIDER', 'keepa')
    result = run(tmp_path)
    assert result['live_requests'] == 0 and result['mock_keepa_requests'] == 8
    assert result['input_rows'] == result['distinct_observations'] == 10
    assert all(not row['purchase_ready'] and row['confidence'] <= 50 for row in result['rows'])
    rows = {row['product']: row for row in result['rows']}
    assert rows['DEMO passing']['status'] == 'REVIEW'
    for name, reason in [('DEMO missing-fees','Missing cost'), ('DEMO package-mismatch','Package quantity'),
                         ('DEMO stale','Stale'), ('PILOT no match','No deterministic'), ('PILOT ambiguous','Ambiguous')]:
        assert rows[name]['status'] == 'REJECT' and reason in rows[name]['reason']
    assert any(e['kind'] == 'amazon_price' and e['value'] is None for e in result['enrichment'])
    assert any(e['kind'] == 'amazon_price' and e['fresh'] is False for e in result['enrichment'])
    assert all(e['value'] is None and e['observed_at'] is None for e in result['enrichment'] if e['kind'] in ('fees','eligibility'))
    with sqlite3.connect(tmp_path/'state/pilot.sqlite3') as db:
        assert db.execute('SELECT count(*) FROM source_observations').fetchone()[0] == 10
        assert db.execute('SELECT count(*) FROM marketplace_matches').fetchone()[0] == 10
        assert db.execute('SELECT count(*) FROM enrichment_observations').fetchone()[0] == 48
        assert db.execute('SELECT count(*) FROM pilot_run_items').fetchone()[0] == 10
        assert db.execute('SELECT status,candidates FROM runs').fetchone() == ('succeeded',10)
        assert db.execute('PRAGMA foreign_key_check').fetchall() == []
        assert db.execute('PRAGMA integrity_check').fetchone() == ('ok',)
        for (payload,) in db.execute('SELECT result_json FROM pilot_run_items'):
            row = json.loads(payload)
            assert row['rejection_reasons'] == rows[row['product']]['rejection_reasons']
        assert json.loads(db.execute('SELECT manifest_json FROM pilot_runs').fetchone()[0]) == result['manifest']
    workbook = load_workbook(tmp_path/'state/PILOT-FIXTURE.xlsx')
    assert len(workbook.sheetnames) == 7 and 'DEMO Enrichment' in workbook.sheetnames
    assert 'FICTIONAL' in workbook.worksheets[0]['A1'].value
    assert not any(cell.data_type=='f' for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    assert json.loads((tmp_path/'state/PILOT-FIXTURE.json').read_text())['label'] == pilot.LABEL
    network.assert_not_called()


def test_duplicate_rows_reimports_keep_run_linkage(tmp_path):
    source = source_rows(tmp_path, count=2, duplicate=True)
    first = run(tmp_path, source)
    second = run(tmp_path, source)
    assert first['run_id'] != second['run_id']
    assert first['input_rows'] == second['input_rows'] == 2
    assert first['distinct_observations'] == second['distinct_observations'] == 1
    assert first['mock_keepa_requests'] == second['mock_keepa_requests'] == 1
    with sqlite3.connect(tmp_path/'state/pilot.sqlite3') as db:
        for table, count in [('source_observations',1),('marketplace_matches',1),
                              ('enrichment_observations',6),('pilot_run_items',4),('pilot_run_enrichment',24)]:
            assert db.execute('SELECT count(*) FROM '+table).fetchone()[0] == count
        assert db.execute('SELECT count(DISTINCT run_id) FROM pilot_run_items').fetchone()[0] == 2


def test_unknowns_not_promoted_from_keepa(tmp_path):
    raw = json.loads((ROOT/'sample_data/pilot_catalog_fixture.json').read_text())
    raw[0].update(seller_eligible=None, sale_price=None, monthly_sales=None, offer_count=None, costs={})
    catalog = tmp_path/'catalog.json'
    catalog.write_text(json.dumps(raw))
    result = pilot.run_offline(source_rows(tmp_path), catalog,
        ROOT/'sample_data/pilot_keepa_fixture.json', tmp_path/'state')
    row, = result['rows']
    assert row['sale_price'] is None and row['monthly_sales'] is None and row['profit'] is None
    for reason in ('Seller eligibility not verified','Missing sale price','Missing demand indicator',
                   'Missing competition indicator','Missing cost referral_fee','Economics cannot be calculated'):
        assert reason in row['rejection_reasons']
    assert any(e['kind']=='bought_past_month' and e['value'] is not None for e in result['enrichment'])
    assert any(e['kind']=='amazon_price' and e['value'] is not None for e in result['enrichment'])


def test_enrichment_failure_keeps_all_screen_rejections_and_continues(tmp_path, monkeypatch):
    monkeypatch.setattr(pilot.FixtureKeepaTransport, 'product', Mock(side_effect=TimeoutError('SECRET')))
    result = run(tmp_path)
    assert len(result['rows']) == 10
    assert all(r['status']=='REJECT' for r in result['rows'])
    assert 'SECRET' not in canonical(result)
    row = next(r for r in result['rows'] if r['product']=='DEMO missing-fees')
    assert 'Keepa enrichment unavailable' in row['rejection_reasons']
    assert 'Missing cost fulfillment_fee' in row['rejection_reasons']


def test_missing_permission_reference_is_preserved_as_rejection(tmp_path):
    row, = run(tmp_path, source_rows(tmp_path, authorization_reference=''))['rows']
    assert 'Missing authorization_reference' in row['rejection_reasons']
    with sqlite3.connect(tmp_path/'state/pilot.sqlite3') as db:
        assert db.execute('SELECT authorization_reference FROM source_observations').fetchone() == ('',)


def test_invalid_normalized_fields_retain_issues_nulls_and_match_reasons(tmp_path):
    source = source_rows(tmp_path, gtin='not-a-valid-code', retrieved_at='not-a-time',
                         variation='different variation', purchase_price='unknown')
    row, = run(tmp_path, source)['rows']
    for reason in ('Invalid gtin', 'Invalid retrieved_at', 'Invalid purchase_price',
                   'Variation mismatch', 'Exact match not verified', 'Missing source timestamp'):
        assert reason in row['rejection_reasons']
    with sqlite3.connect(tmp_path/'state/pilot.sqlite3') as db:
        source_payload = json.loads(db.execute('SELECT payload_json FROM source_observations').fetchone()[0])
        match_payload = json.loads(db.execute('SELECT payload_json FROM marketplace_matches').fetchone()[0])
        assert source_payload['purchase_price'] is None and source_payload['retrieved_at'] is None
        assert 'Invalid gtin' in source_payload['issues']
        assert 'Variation mismatch' in match_payload['reasons']
        assert db.execute('SELECT observed_at,purchase_price FROM source_observations').fetchone() == (None,None)


@pytest.mark.parametrize('fault', ['empty', 'limit', 'report'])
def test_failed_runs_are_auditable(tmp_path, monkeypatch, fault):
    source = source_rows(tmp_path, count=0 if fault=='empty' else 10)
    if fault=='report':
        monkeypatch.setattr(pilot, 'write_report', Mock(side_effect=OSError('SECRET')))
    with pytest.raises((ValueError,OSError)):
        run(tmp_path, source, maximum_products=1 if fault=='limit' else 10)
    with sqlite3.connect(tmp_path/'state/pilot.sqlite3') as db:
        assert db.execute('SELECT status FROM runs').fetchone() == ('failed',)
        assert db.execute('PRAGMA foreign_key_check').fetchall() == []
    assert not (tmp_path/'state/PILOT-FIXTURE.json').exists()


def test_fixture_label_required_and_input_output_collision(tmp_path):
    keepa = tmp_path/'not-fixture.json'
    keepa.write_text('{"responses":{}}')
    with pytest.raises(ValueError):
        pilot.FixtureKeepaTransport(keepa)
    directory = tmp_path/'collision'
    directory.mkdir()
    source = directory/'PILOT-FIXTURE.json'
    source.write_text('preserve me')
    with pytest.raises(ValueError):
        pilot.run_offline(source,
            ROOT/'sample_data/pilot_catalog_fixture.json', ROOT/'sample_data/pilot_keepa_fixture.json',directory)
    assert source.read_text() == 'preserve me'


def test_storage_precision_changes_and_transactional_linkage(tmp_path):
    store = ObservationStore(tmp_path/'evidence.db')
    product = next(CsvConnector(ROOT/'sample_data/pilot_source_fixture.csv').products())
    match = MockAmazonProvider(ROOT/'sample_data/pilot_catalog_fixture.json').match(product)
    run_id = store.start_run(pilot.FIXTURE_NOW, {'fixture':True})
    evidence = Evidence(None, pilot.FIXTURE_NOW, 'FICTIONAL unknown', True)
    payload = [('B000000001','fees',evidence,pilot.FIXTURE_NOW)]
    row = dict(status='REJECT',rejection_reasons=['Unknown fees'])
    with pytest.raises(sqlite3.IntegrityError):
        store.record_item(99999,1,product,match,row,payload,pilot.FIXTURE_NOW)
    with store.connect() as db:
        assert db.execute('SELECT count(*) FROM source_observations').fetchone() == (0,)
    store.record_item(run_id,1,product,match,row,payload,pilot.FIXTURE_NOW)
    for ordinal, changed in enumerate((replace(product,purchase_price=Decimal('11.1234567890123456789012345678901')),
                                        replace(product,retrieved_at=product.retrieved_at+timedelta(minutes=1)),
                                        replace(product,authorization_reference='different license')),2):
        store.record_item(run_id,ordinal,changed,match,row,payload,pilot.FIXTURE_NOW)
    with store.connect() as db:
        assert db.execute('SELECT count(*) FROM source_observations').fetchone() == (4,)
        assert db.execute('SELECT purchase_price FROM source_observations WHERE id=2').fetchone()[0] == '11.1234567890123456789012345678901'
        assert db.execute('SELECT observed_at FROM enrichment_observations').fetchone() == (None,)
    assert fingerprint(replace(product,purchase_price=Decimal('10.00'))) == fingerprint(product)
    offset = timezone(timedelta(hours=-5))
    assert fingerprint(replace(product,retrieved_at=product.retrieved_at.astimezone(offset))) == fingerprint(product)
    with pytest.raises(ValueError):
        canonical(dict(time=pilot.FIXTURE_NOW.replace(tzinfo=None)))
    assert store.path.stat().st_mode & 0o777 == 0o600
    ObservationStore(store.path)
    with store.connect() as db:
        db.execute('UPDATE evidence_schema SET version=99')
    with pytest.raises(ValueError):
        ObservationStore(store.path)


def test_private_database_and_symlink_rejection(tmp_path):
    target = tmp_path/'target.db'
    ObservationStore(target)
    link = tmp_path/'link.db'
    link.symlink_to(target)
    with pytest.raises(ValueError):
        ObservationStore(link)
    target.chmod(0o644)
    with pytest.raises(ValueError):
        ObservationStore(target)
