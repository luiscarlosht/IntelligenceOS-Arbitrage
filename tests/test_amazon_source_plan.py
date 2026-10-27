from dataclasses import replace
from datetime import timedelta
from decimal import Decimal
import sqlite3

import pytest
from openpyxl import load_workbook

from app.amazon_source_plan import prepare, RestrictionCheck
from app.amazon_details import PricingSnapshot
from app.amazon_offer_coverage import offer_coverage
from app.config import ROOT
from app.marketplace import MockAmazonProvider
from app.offline_pilot import FIXTURE_NOW
from app.provider_contracts import Evidence


def inputs():
    listings = MockAmazonProvider(ROOT/'sample_data/pilot_catalog_fixture.json').listings
    pricing = {l.asin: Evidence(PricingSnapshot(l.asin, 'ATVPDKIKX0DER', 'New', (), 1,
                price_unknown=False, exclude_from_profitability_ranking=False), FIXTURE_NOW, 'synthetic', True)
               for l in listings}
    restrictions = {l.asin: Evidence(RestrictionCheck(l.asin, 'owner', 'ATVPDKIKX0DER', 'new_new', True),
                    FIXTURE_NOW, 'synthetic', True) for l in listings}
    return listings, pricing, restrictions


def run(tmp_path, listings, pricing, restrictions, prices=None):
    return prepare(ROOT/'sample_data/pilot_source_fixture.csv', listings, pricing, restrictions,
                   prices or {1:'37.25'}, tmp_path, seller_reference='owner', now=FIXTURE_NOW,
                   fba_eligibility={l.asin: Evidence(True, FIXTURE_NOW, 'synthetic FBA eligibility', True) for l in listings})


def test_plan_persists_all_ten_and_only_explicit_fee_scenario(tmp_path, monkeypatch):
    def forbidden(*a, **k):
        raise AssertionError('No external action authorized')
    monkeypatch.setattr('urllib.request.OpenerDirector.open', forbidden)
    monkeypatch.setattr('app.mailer.send_report', forbidden)
    result = run(tmp_path, *inputs())
    assert len(result['rows']) == len(result['scenarios']) == 10
    assert result['scenarios'][0]['fee_request_planned']
    assert result['scenarios'][0]['listing_price_usd'] == Decimal('37.25')
    assert not any(s['fee_request_planned'] for s in result['scenarios'][1:])
    assert result['budget']['maximum_attempts'] == 32
    assert result['budget']['live_requests'] == 0
    with sqlite3.connect(tmp_path/'source-plan.sqlite3') as db:
        assert db.execute('SELECT count(*) FROM pilot_run_items').fetchone()[0] == 10
        assert db.execute('SELECT status FROM runs').fetchone()[0] == 'succeeded'
        assert db.execute('PRAGMA foreign_key_check').fetchall() == []
    workbook = load_workbook(tmp_path/'SOURCE-PLAN.xlsx')
    assert 'All Candidates' in workbook.sheetnames


@pytest.mark.parametrize('failure', ['restricted', 'scope', 'partial', 'stale', 'missing', 'eligibility'])
def test_unsuitable_candidates_keep_reasons_and_skip_fees(tmp_path, failure):
    listings, pricing, restrictions = inputs()
    asin = listings[0].asin
    if failure == 'restricted':
        restrictions[asin] = replace(restrictions[asin], value=replace(restrictions[asin].value,
                                     listing_permitted=False, reason_codes=('NOT_ELIGIBLE',)))
    elif failure == 'scope':
        restrictions[asin] = replace(restrictions[asin], value=replace(restrictions[asin].value, seller_reference='other'))
    elif failure == 'partial':
        pricing[asin] = replace(pricing[asin], value=replace(pricing[asin].value, exclude_from_profitability_ranking=True))
    elif failure == 'stale':
        pricing[asin] = replace(pricing[asin], observed_at=FIXTURE_NOW-timedelta(days=2))
    elif failure == 'missing':
        pricing.clear()
    else:
        listings = (replace(listings[0], seller_eligible=None), *listings[1:])
    result = run(tmp_path, listings, pricing, restrictions)
    assert not result['scenarios'][0]['fee_request_planned']
    assert result['scenarios'][0]['reasons']
    assert result['rows'][0]['status'] == 'REJECT'
    assert len(result['rows']) == 10


def test_retained_asin_is_excluded_even_if_supplied_catalog_claims_eligible(tmp_path):
    listings, pricing, restrictions = inputs()
    listings = (replace(listings[0], asin='B0G1MJ24B4'), *listings[1:])
    # Existing deterministic matching still rejects hints that conflict with identity.
    from app.connectors import CsvConnector
    from app.models import Match
    from app.discovery import screen
    from app.config import Config
    product = next(CsvConnector(ROOT/'sample_data/pilot_source_fixture.csv').products())
    row = screen(product, Match(listings[0], True, Decimal('1')), Config(
                 input=tmp_path/'i', output=tmp_path/'o', database=tmp_path/'d'), FIXTURE_NOW)
    assert row['status'] == 'REJECT' and 'NOT_ELIGIBLE' in row['reason']
    assert row['ranking_score'] == 0


def test_condition_breakdown_explains_count_without_clearing_aggregate_warning():
    summary = {'NumberOfOffers': [dict(condition='new', fulfillmentChannel='Amazon', OfferCount=1),
                                  dict(condition='used', fulfillmentChannel='Amazon', OfferCount=4)]}
    detail = dict(SubCondition='new', IsFulfilledByAmazon=True, ShippingTime={},
                  ListingPrice=dict(CurrencyCode='USD', Amount=Decimal('8.97')),
                  Shipping=dict(CurrencyCode='USD', Amount=Decimal('0')))
    result = offer_coverage(summary, [detail], 5)
    assert result.requested_condition_offer_count == 1 and result.condition_count_matches_returned
    assert result.coverage == 'partial' and result.incomplete_coverage
    assert result.documented_expected_offer_count == 5 and not result.documented_count_matches
    summary['NumberOfOffers'] = [dict(OfferCount=5)]
    result = offer_coverage(summary, [detail], 5)
    assert result.requested_condition_offer_count is None
    assert result.condition_count_matches_returned is None


def test_missing_independent_fba_eligibility_blocks_fee_plan(tmp_path):
    listings, pricing, restrictions = inputs()
    result = prepare(ROOT/'sample_data/pilot_source_fixture.csv', listings, pricing, restrictions,
                     {1:'37.25'}, tmp_path, seller_reference='owner', now=FIXTURE_NOW)
    assert not result['scenarios'][0]['fee_request_planned']
    assert 'Missing or unverified FBA eligibility' in result['scenarios'][0]['reasons']


@pytest.mark.parametrize('prices', [{0:'12.50'}, {1:'0'}, {1:'NaN'}, {1:'12.345'}])
def test_invalid_explicit_scenarios_refused_before_persistence(tmp_path, prices):
    with pytest.raises(ValueError):
        run(tmp_path, *inputs(), prices=prices)
    assert not (tmp_path/'source-plan.sqlite3').exists()


def test_initial_ingestion_needs_no_amazon_fees_or_eligibility(tmp_path):
    result = prepare(ROOT/'sample_data/pilot_source_fixture.csv', (), {}, {}, {},
                     tmp_path, seller_reference='owner', now=FIXTURE_NOW)
    assert len(result['rows']) == 10
    assert all(row['status'] == 'REJECT' for row in result['rows'])
    assert all(not item['fee_request_planned'] for item in result['scenarios'])
    assert all(row['sale_price'] is None and row['profit'] is None for row in result['rows'])
    assert all('Missing marketplace listing' in row['rejection_reasons'] for row in result['rows'])
    assert result['budget']['fees'] == result['budget']['live_requests'] == 0
    with sqlite3.connect(tmp_path/'source-plan.sqlite3') as db:
        assert db.execute('SELECT count(*) FROM pilot_run_items').fetchone()[0] == 10
