"""Synthetic transports only; source rows remain unique while candidates stay independent."""
from copy import deepcopy
import json
import sqlite3
from urllib.parse import urlsplit

import pytest

from app import candidate_execution as executor
from app.observations import canonical
from test_candidate_execution import stored_case
from test_amazon_candidates import catalog, CREDS


@pytest.mark.parametrize('failure_at', [None, 1, 3])
def test_ties_shared_requests_failure_and_replay(tmp_path, monkeypatch, failure_at):
    _, _, _, path, _, manifest = stored_case(tmp_path, count=2, shared_asin=True)
    manifest.update(mode='reviewed_shortlist_details', private_directory=str(tmp_path/'private'))
    second = 'B004U42A28'
    for row in manifest['rows']:
        first = dict(candidate_asin=row['candidate_asin'], catalog_detail_reuse=[],
                     selection_reason='Tie retained', candidate_audit={}, alternatives=[])
        row['candidate_outcomes'] = [first, dict(deepcopy(first), candidate_asin=second)]
        row['search']['selected_asin'] = None
    request = deepcopy(manifest['requests'][0])
    from app.observations import fingerprint
    from app.candidate_execution import MARKETPLACE, DETAIL_DATA
    request.update(asin=second, path='/catalog/2022-04-01/items/'+second,
        request_key=fingerprint(dict(kind='catalog_detail', marketplace_id=MARKETPLACE,
                                    asin=second, included_data=DETAIL_DATA)))
    manifest['requests'].append(request)
    manifest['limits'].update(detail=2, total=3)
    path.write_text(canonical(manifest))
    import hashlib
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(executor, 'prepare_shortlist', lambda _: deepcopy(manifest))
    calls = []
    def transport(req):
        calls.append(req)
        if len(calls) == failure_at:
            raise RuntimeError('synthetic failure')
        if len(calls) == 1:
            return dict(access_token='synthetic-access', token_type='bearer')
        return catalog(urlsplit(req.full_url).path.rsplit('/',1)[-1])
    report = executor.execute_details(path, digest, 3, tmp_path/'private',
        credential_loader=lambda _: CREDS, transport=transport, sleep=lambda _: None)
    assert len(report['rows']) == 2
    assert all(len(r['candidate_outcomes']) == 2 for r in report['rows'])
    assert all(r['search']['selected_asin'] is None for r in report['rows'])
    assert all(not c['exact'] for r in report['rows'] for c in r['candidate_outcomes'])
    assert len(calls) == (failure_at or 3)
    if failure_at == 3:
        assert all([c['status'] for c in r['candidate_outcomes']] == ['candidate_unverified', 'failed']
                   for r in report['rows'])
    if failure_at is None:
        for row in report['rows']:
            assert [c['detail']['value']['catalog']['asin'] for c in row['candidate_outcomes']] == [
                manifest['requests'][0]['asin'], second]
    with pytest.raises(sqlite3.IntegrityError):
        executor.execute_details(path, digest, 3, tmp_path/'private',
            credential_loader=lambda _: pytest.fail('Replay credentials'),
            transport=lambda _: pytest.fail('Replay HTTP'))
    from openpyxl import load_workbook
    workbook = load_workbook(next((tmp_path/'private').glob('*.xlsx')))
    assert workbook['Reviewed Candidate Outcomes'].max_row == 5
    workbook.close()


def test_nested_saved_evidence_reuse_retains_original_timestamp(tmp_path):
    from app.catalog_saved_evidence import saved_evidence
    from app.observations import normalized, fingerprint
    from app.amazon_candidates import detail_evidence
    from app.candidate_execution import MARKETPLACE
    asin = 'B004U42A28'
    source = dict(retailer='Walmart', source_context={})
    detail = normalized(detail_evidence(catalog(asin), asin, is_mock=False, secrets=()))
    report = dict(is_mock=False, manifest=dict(mode='reviewed_shortlist_details', requests=[
        dict(kind='catalog_detail', asin=asin, source_record_ids=['r1'],
             params=dict(marketplaceIds=MARKETPLACE))]),
        rows=[dict(ordinal=1, record_id='r1', source=source, source_fingerprint=fingerprint(source),
                   candidate_outcomes=[dict(search=dict(selected_asin=asin), detail=detail)])],
        attempts=[dict(kind='detail', asin=asin, status='response_received')])
    path = tmp_path/'candidate-details-nested.json'
    path.write_text(canonical(report))
    cache, blocked, _ = saved_evidence([path])
    assert cache[asin]['observed_at'] == detail['observed_at'] and asin in blocked
    report['attempts'][0]['status'] = 'reserved'
    path.write_text(canonical(report))
    cache, blocked, _ = saved_evidence([path])
    assert cache == {} and asin in blocked
