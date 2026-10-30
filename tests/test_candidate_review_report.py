"""Saved-evidence presentation tests only; no network or credential access."""
from copy import deepcopy
from decimal import Decimal
import hashlib
import json

import pytest
from openpyxl import load_workbook

from app.candidate_review_report import build_review, evidence_gaps, regenerate
from app.reporting import write_report


def saved_fixture(tmp_path):
    items=[dict(name=f'Synthetic utensil {i}',url=f'https://www.ebay.com/itm/{100000000000+i}?tracking=original',
                shop_name='synthetic-seller',price='12.99',shipping_costs=None if i==0 else '0.0') for i in range(10)]
    export=dict(status='finished',results=[dict(success=True,query=dict(source='ebay',country='us',topic='search_results'),
                metadata=dict(page_count_live=1),content=dict(search_results=items))])
    source_file=tmp_path/'source.json';source_file.write_text(json.dumps(export))
    digest=hashlib.sha256(source_file.read_bytes()).hexdigest()
    rows=[]
    for i,item in enumerate(items,1):
        asin=f'B{i:09d}'
        source=dict(retailer='eBay',name=item['name'],url=item['url'],gtin='',manufacturer='',model='',
            package_quantity=None,variation='',retrieved_at=None,availability='unknown',
            provenance='Synthetic saved Price API export; sha256='+digest,
            source_context=dict(source_seller=item['shop_name'],search_price=item['price'],
                source_shipping_cost=item['shipping_costs'],source_package_contents=None))
        catalog=dict(asin=asin,title=f'Catalog utensil {i}',gtin='',manufacturer='Maker',model='M1',
                     package_quantity=1,variation='{"color":"red"}',is_mock=False,mapping_issues=[])
        rows.append(dict(ordinal=i,status='ambiguous',exact=False,source=source,
            search=dict(selected_asin=asin,candidates=[dict(asin=asin,title=catalog['title'])],complete=False,more_results=True),
            detail=dict(is_mock=False,observed_at='2026-10-27T23:10:00+00:00',provenance='Saved LIVE catalog',
                        value=dict(catalog=catalog,package_attributes=dict(number_of_items=['10'])))))
    report=dict(status='succeeded',is_mock=False,live_requests=21,attempts=[{}]*21,rows=rows)
    report_file=tmp_path/'report.json';report_file.write_text(json.dumps(report))
    old=tmp_path/'original.xlsx';old.write_bytes(b'synthetic original sentinel')
    return report,report_file,source_file,old


def test_live_report_comparison_unknown_economics_and_preserved_inputs(tmp_path,monkeypatch):
    report,rf,sf,old=saved_fixture(tmp_path)
    before={p:p.read_bytes() for p in (rf,sf,old)}
    monkeypatch.setattr('urllib.request.OpenerDirector.open',lambda *a,**k:pytest.fail('Network'))
    output=tmp_path/'corrected.xlsx'
    result=regenerate(rf,sf,old,output)
    assert result['rows']==10 and result['candidate_only']==10 and result['confirmed']==result['unmatched']==0
    assert all(p.read_bytes()==b for p,b in before.items())
    wb=load_workbook(output)
    assert 'LIVE eBay search' in wb['Executive Summary']['B8'].value
    assert 'LIVE Amazon' in wb['Executive Summary']['B9'].value
    assert 'UNKNOWN' in wb['Executive Summary']['B10'].value
    assert not any(isinstance(c.value,str) and ('MOCK' in c.value or 'Mock data' in c.value)
                   for sheet in wb for row in sheet for c in row)
    comparison=wb['Candidate Comparison']
    assert comparison.max_row==11 and comparison.freeze_panes=='C2' and comparison.auto_filter.ref==comparison.dimensions
    assert comparison['C2'].value==12.99
    assert comparison['D2'].value is None and comparison['D3'].value==0
    assert comparison['E2'].hyperlink.target==report['rows'][0]['source']['url']
    assert comparison['E2'].value=='https://www.ebay.com/itm/100000000000'
    assert comparison['F2'].value=='B000000001' and comparison['G2'].value=='Catalog utensil 1'
    assert comparison.column_dimensions['I'].width==85
    assert 'Source retail package count unknown' in comparison['I2'].value
    assert 'CONFLICT' not in comparison['I2'].value
    for name in ('All Candidates','Rejected Products'):
        sheet=wb[name];headers={c.value:c.column for c in sheet[1]}
        for row in range(2,12):
            for key in ('Buy Cost','Sale Price','Profit','Roi','Acquisition Cost','Referral Fee','Fulfillment Fee'):
                assert sheet.cell(row,headers[key]).value is None
            assert sheet.cell(row,headers['Purchase Ready']).value is False
            assert 'LIVE Amazon' in sheet.cell(row,headers['Data Mode']).value
            assert sheet.cell(row,headers['Amazon Asin']).value is not None
    wb.close()
    with pytest.raises(ValueError):regenerate(rf,sf,old,output)
    with pytest.raises(ValueError):regenerate(rf,sf,old,old)


def test_formula_safety_in_comparison_and_numeric_prices(tmp_path):
    report,*_=saved_fixture(tmp_path)
    row=report['rows'][0];row['source']['name']='=HYPERLINK("evil")';row['source']['source_context']['source_seller']='=1+1'
    row['detail']['value']['catalog']['title']='@SUM(A1:A2)'
    rows,review,evidence=build_review(report)
    output=tmp_path/'formula.xlsx'
    write_report(rows,output,discovery=True,candidate_review=review,enrichment_rows=evidence)
    wb=load_workbook(output)
    assert all(wb['Candidate Comparison'][address].data_type=='s' for address in ('A2','B2','G2'))
    assert not any(c.data_type=='f' for s in wb for r in s for c in r)
    wb.close()


def test_missing_and_demonstrated_conflicts_are_separate(tmp_path):
    report,*_=saved_fixture(tmp_path);row=report['rows'][0]
    source=row['source'];catalog=row['detail']['value']['catalog'];search=row['search']
    assert not any('CONFLICT' in gap for gap in evidence_gaps(source,catalog,search))
    source.update(gtin='036000291452',manufacturer='Other maker',model='Other model',package_quantity=2,
                  variation='{"color":"blue"}')
    catalog['gtin']='4006381333931'
    gaps=evidence_gaps(source,catalog,search)
    for field in ('GTIN','manufacturer','model','retail package quantity','variation color'):
        assert any(gap.startswith('CONFLICT:') and field in gap for gap in gaps)


def test_mock_evidence_keeps_mock_labels_and_cannot_regenerate_as_live(tmp_path):
    report,rf,sf,old=saved_fixture(tmp_path);report['is_mock']=True
    for row in report['rows']:
        row['detail']['is_mock']=row['detail']['value']['catalog']['is_mock']=True
    rows,review,enrichments=build_review(report)
    assert 'MOCK' in review['Amazon evidence'] and all('MOCK' in r['data_mode'] for r in rows)
    rf.write_text(json.dumps(report))
    with pytest.raises(ValueError):regenerate(rf,sf,old,tmp_path/'bad.xlsx')
    with pytest.raises(ValueError):write_report(rows,tmp_path/'demo.xlsx',discovery=True,demo=True,candidate_review=review)


@pytest.mark.parametrize('change',['detail_asin','mixed_mock','export_tamper','source_price','source_shipping','ordinals'])
def test_evidence_scope_tampering_rejected(tmp_path,change):
    report,rf,sf,old=saved_fixture(tmp_path)
    if change=='detail_asin':report['rows'][0]['detail']['value']['catalog']['asin']='B099999999'
    if change=='mixed_mock':report['rows'][0]['detail']['is_mock']=True
    if change=='export_tamper':sf.write_text(sf.read_text()+' ')
    if change=='source_price':report['rows'][0]['source']['source_context']['search_price']='999'
    if change=='source_shipping':report['rows'][0]['source']['source_context']['source_shipping_cost']='99'
    if change=='ordinals':report['rows'][0]['ordinal']=2
    rf.write_text(json.dumps(report))
    with pytest.raises(ValueError):regenerate(rf,sf,old,tmp_path/'bad.xlsx')
    assert not (tmp_path/'bad.xlsx').exists()
