"""Separate one-shot CORRECTED candidate-review delivery; interactive SSH only, no retries."""
import argparse,hashlib,json,os,re,resource,shlex,signal,sys
from pathlib import Path
from datetime import datetime,timezone
from email.message import EmailMessage
from email.utils import formatdate,make_msgid
from io import BytesIO
from decimal import Decimal
from openpyxl import load_workbook
from app.gmail_api import send_message
from ops.gmail_stage_tracking import StageTracker

root=Path(__file__).resolve().parents[1]
SENDER_CONFIG=Path('/root/.bashrc')
base=Path('/root/.intelligenceos-arbitrage-amazon')
stem='candidates-cabbbda426bff05faff0b6d8f051e62bde21fa927efafe9fc347e9380ee24f11'
report=base/(stem+'.json'); workbook=base/(stem+'-corrected.xlsx')
marker=root/'var/diagnostics/gmail-ebay-candidate-review-corrected-ssh-1.attempted'
stages=root/'var/diagnostics/gmail-ebay-candidate-review-corrected-ssh-1.stages.jsonl'
outcome_path=root/'var/diagnostics/gmail-ebay-candidate-review-corrected-ssh-1.outcome.json'
EXPECTED_WORKBOOK_SHA256='23138171a64294d9e3353a5cfecd32d303030ca181ff59dbc8cbd0fac396dd47'

def check_report():
    data=json.loads(report.read_text())
    if not (data['status']=='succeeded' and data['live_requests']==21 and len(data['attempts'])==21):
        raise ValueError('Candidate report or message verification failed')
    if not (not data['is_mock'] and len(data['rows'])==10):
        raise ValueError('Candidate report or message verification failed')
    if not (all(r['status']=='ambiguous' and r['exact'] is False for r in data['rows'])):
        raise ValueError('Candidate report or message verification failed')
    if not (workbook.is_file() and not workbook.is_symlink()):
        raise ValueError('Candidate report or message verification failed')
    content=workbook.read_bytes();digest=hashlib.sha256(content).hexdigest()
    if not (digest==EXPECTED_WORKBOOK_SHA256):
        raise ValueError('Candidate report or message verification failed')
    wb=load_workbook(BytesIO(content),read_only=True,data_only=False)
    try:
        comparison=wb['Candidate Comparison']
        summary=dict(wb['Executive Summary'].values)
        if (comparison.max_row!=11 or summary.get('Candidates')!=10 or
                summary.get('Confirmed exact matches')!=0 or summary.get('Candidate-only rows')!=10 or
                'LIVE eBay search observations' not in str(summary.get('Source evidence')) or
                'LIVE Amazon catalog observations' not in str(summary.get('Amazon evidence')) or
                'UNKNOWN' not in str(summary.get('Economics'))):
            raise ValueError('Corrected workbook provenance/summary mismatch')
        if any(isinstance(c.value,str) and ('MOCK' in c.value or 'Mock data' in c.value)
               for sheet in wb for row in sheet for c in row):
            raise ValueError('Incorrect mock provenance label')
        for i,result in enumerate(data['rows'],2):
            source=result['source'];context=source['source_context'];catalog=result['detail']['value']['catalog']
            if (comparison.cell(i,1).value!=source['name'] or comparison.cell(i,2).value!=context['source_seller'] or
                    comparison.cell(i,6).value!=catalog['asin'] or comparison.cell(i,7).value!=catalog['title'] or
                    comparison.cell(i,8).value!='Ambiguous candidate only'):
                raise ValueError('Corrected comparison differs from saved evidence')
            for col,key in ((3,'search_price'),(4,'source_shipping_cost')):
                actual=comparison.cell(i,col).value;expected=context.get(key)
                if (actual is None)!=(expected is None):
                    raise ValueError('Observed price/shipping unknown mismatch')
                if actual is not None and Decimal(str(actual))!=Decimal(str(expected)):
                    raise ValueError('Observed price/shipping value mismatch')
        headers={cell.value:cell.column for cell in wb['All Candidates'][1]}
        for row in range(2,12):
            for field in ('Buy Cost','Sale Price','Profit','Roi','Acquisition Cost','Referral Fee','Fulfillment Fee'):
                if wb['All Candidates'].cell(row,headers[field]).value is not None:
                    raise ValueError('Unsupported economics present')
    finally:
        wb.close()
    return content,digest


def execute():
    content,digest=check_report()
    addresses={}
    # Select the existing nonsecret sender assignment without loading SMTP passwords
    # or executing shell configuration. Delivery uses only the existing OAuth loader.
    for line in SENDER_CONFIG.open():
        match=re.match(r'^\s*(?:export\s+)?(EMAIL_FROM|SMTP_USER|SMTP_USERNAME)=(.*)$',line.rstrip('\n'))
        if match:
            words=shlex.split(match[2],comments=True)
            if len(words)==1:addresses[match[1]]=words[0]
    senders=set(addresses.values())
    if not (len(senders)==1):
        raise ValueError('Candidate report or message verification failed')
    sender=next(iter(senders))
    if not (re.fullmatch(r'[A-Za-z0-9._%+\-]+@gmail\.com',sender,re.I)):
        raise ValueError('Candidate report or message verification failed')
    # A permanent exclusive reservation precedes OAuth refresh or message submission.
    if stages.exists() or outcome_path.exists():
        raise FileExistsError('Attempt evidence already exists')
    fd=os.open(marker,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
    with os.fdopen(fd,'w') as f:
        json.dump(dict(reserved_at=datetime.now(timezone.utc).isoformat(),
                       purpose='Corrected — eBay/Amazon proof-of-concept candidate review',workbook_sha256=digest,
                       automatic_retry=False,maximum_messages=1),f)
        f.flush();os.fsync(f.fileno())
    result_fd=os.open(outcome_path,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
    os.close(os.open(stages,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600))
    tracker=StageTracker(stages)
    result=dict(status='reserved',acceptance_observed=False,submission_started=False,automatic_retry=False)
    def persist():
        os.lseek(result_fd,0,os.SEEK_SET)
        payload=(json.dumps(result)+'\n').encode();os.write(result_fd,payload);os.ftruncate(result_fd,len(payload));os.fsync(result_fd)
    persist()
    def deadline(*_):raise TimeoutError()
    previous=signal.signal(signal.SIGALRM,deadline);signal.alarm(90)
    try:
        tracker('message_construction','started')
        message=EmailMessage()
        message['From']=sender
        message['To']='luiscarloshernandez@hotmail.com'
        message['Subject']='Corrected — eBay/Amazon proof-of-concept candidate review'
        message['Date']=formatdate(localtime=False)
        message['Message-ID']=make_msgid(domain='intelligenceos.local')
        message.set_content('Corrected — eBay/Amazon proof-of-concept candidate review\n\n'
            'Attached is the corrected real eBay-to-Amazon candidate-review workbook from the successful 21-request Amazon catalog run.\n\n'
            '• 10 real eBay listings reviewed.\n'
            '• 10 ambiguous candidates; zero confirmed matches.\n'
            '• Profitability unknown.\n\n'
            'Successful API execution does not verify product identity. Source identifiers, package contents/count and variation evidence still require validation.\n\n'
            'This is a proof of concept for candidate review, not a purchase recommendation.\n')
        message.add_attachment(content,maintype='application',
            subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='Corrected-eBay-Amazon-Candidate-Review.xlsx')
        if not (len(list(message.iter_attachments()))==1):
            raise ValueError('Candidate report or message verification failed')
        tracker('message_construction','completed')
        send_message(message,Path('/root/.config/arbitrage-gmail/token.json'),observe=tracker)
        result['status']='gmail_accepted' if tracker.acceptance_observed else 'acceptance_not_observed'
    except (Exception,KeyboardInterrupt):
        result['status']='submission_uncertain' if tracker.submission_started else 'failed_before_submission'
    finally:
        signal.alarm(0);signal.signal(signal.SIGALRM,previous)
        result.update(acceptance_observed=tracker.acceptance_observed,submission_started=tracker.submission_started,
                      last_completed_stage=tracker.last_completed,finished_at=datetime.now(timezone.utc).isoformat(),
                      report_preserved=hashlib.sha256(workbook.read_bytes()).hexdigest()==digest)
        persist();os.close(result_fd)
    print(json.dumps(result))
    return 0 if result['acceptance_observed'] else 1

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    modes=parser.add_mutually_exclusive_group(required=True)
    modes.add_argument('--check',action='store_true',help='Offline verification; no credentials or send')
    modes.add_argument('--send',action='store_true',help='One fresh attempt, only from an interactive terminal')
    args=parser.parse_args()
    try:
        if args.check:
            _,digest=check_report()
            if any(p.exists() for p in (marker,stages,outcome_path)):
                raise FileExistsError()
            print(json.dumps(dict(status='verified_offline',workbook_sha256=digest,
                                  recipient='luiscarloshernandez@hotmail.com',attempt='corrected-ssh-1',
                                  maximum_messages=1,automatic_retry=False,credentials_read=False)))
            return 0
        if not sys.stdin.isatty():
            raise ValueError('Interactive SSH terminal required')
        resource.setrlimit(resource.RLIMIT_CORE,(0,0))
        return execute()
    except (Exception,KeyboardInterrupt):
        print(json.dumps(dict(status='stopped_without_retry',automatic_retry=False)))
        return 1


if __name__=='__main__':
    raise SystemExit(main())
