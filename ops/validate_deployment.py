"""Read-only validation of bundled offline reports, history and freshness gates.

Run from the deployed release with its venv Python, as the application user.
"""
from datetime import datetime, timedelta, timezone
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

from app.config import Config, ROOT
from app.connectors import CsvConnector
from app.discovery import discover
from app.marketplace import MockAmazonProvider


def main() -> None:
    state = Path('/var/lib/intelligenceos-arbitrage')
    for filename, sheets, reviews in (
        ('legacy-illustrative.xlsx', 4, 1),
        ('discovery-mock.xlsx', 6, None),
    ):
        workbook = load_workbook(state / filename, read_only=True)
        assert len(workbook.sheetnames) == sheets
        rows = list(workbook['All Candidates'].values)
        records = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert len(records) == 3
        if reviews is not None:
            assert sum(row['Status'] == 'REVIEW' for row in records) == reviews
        else:
            assert all('MOCK' in row['Data Mode'] for row in records)
            assert all(row['Purchase Ready'] is False for row in records)
        assert not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row)
        print(filename, workbook.sheetnames, [row['Status'] for row in records])
        workbook.close()

    with sqlite3.connect(f'file:{state / "history.sqlite3"}?mode=ro', uri=True) as db:
        assert db.execute('PRAGMA integrity_check').fetchone() == ('ok',)
        runs = db.execute('SELECT id, status, candidates, review_candidates, report_path, started_at, finished_at FROM runs ORDER BY id DESC LIMIT 2').fetchall()
        assert len(runs) == 2
        assert {Path(row[4]).name for row in runs} == {'legacy-illustrative.xlsx', 'discovery-mock.xlsx'}
        for row in runs:
            assert row[1] == 'succeeded' and row[2] == 3
            assert datetime.fromisoformat(row[5]).utcoffset() is not None
            assert datetime.fromisoformat(row[6]).utcoffset() is not None
        print('SQLite integrity ok; latest runs:', runs)

    config = Config.from_env()
    from dataclasses import replace
    config = replace(config, mode='discovery', input=ROOT / 'sample_data/discovery.csv')
    stale = discover(CsvConnector(config.input), MockAmazonProvider(config.amazon_fixture),
                     config, datetime.now(timezone.utc) + timedelta(days=2))
    assert len(stale) == 3 and all(row['status'] == 'REJECT' for row in stale)
    assert any('stale' in row['reason'].lower() for row in stale)
    print('Future-clock stale fixture check: all 3 rejected; 24-hour gate unchanged')


if __name__ == '__main__':
    main()
